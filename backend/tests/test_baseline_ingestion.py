import base64
import io
import json
from datetime import date
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from app.db.base import Base
from app.db.models import BalanceCheckpoint, Entity, ImportBatch, LedgerAccount, Organization
from app.ingestion.baseline import (
    normalize_balance_checkpoint_xlsx,
    parse_baseline_xlsx,
)
from app.ingestion.batches import activate_import_batch, stage_import_batch


HEADERS = ["Account Code", "Balance Date", "Signed Balance", "Currency"]
BALANCE_DATE = date(2025, 3, 31)


def _xlsx_bytes(rows, headers=HEADERS):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.fixture
def baseline_session():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        organization = Organization(name="Baseline Test Org")
        session.add(organization)
        session.flush()
        entity = Entity(
            organization_id=organization.id,
            name="Baseline Test Entity",
            materiality_threshold=Decimal("0.00"),
        )
        session.add(entity)
        session.flush()
        session.add_all([
            LedgerAccount(entity_id=entity.id, external_code="1000", name="Cash"),
            LedgerAccount(entity_id=entity.id, external_code="2000", name="Capital"),
            LedgerAccount(entity_id=entity.id, external_code="3000", name="Revenue"),
        ])
        session.commit()
        yield session, entity.id
    Base.metadata.drop_all(engine)


def _stage(session, entity_id, contents):
    return stage_import_batch(
        session,
        entity_id=entity_id,
        period_start=date(2025, 4, 1),
        period_end=date(2026, 3, 31),
        source="gl_upload",
        source_family="gl_upload",
        original_filename="closing-baseline.xlsx",
        contents=contents,
        uploaded_by_user_id=None,
        kind="balance_checkpoint",
        coverage_start=date(2025, 4, 1),
        coverage_end=date(2026, 3, 31),
    )


def _valid_bytes():
    return _xlsx_bytes([
        ["1000", BALANCE_DATE, "100.00", "INR"],
        ["2000", BALANCE_DATE, "-60.00", "INR"],
        ["3000", BALANCE_DATE, "-40.00", "INR"],
    ])


def _checkpoint_snapshot(session, batch_id):
    return [
        (
            checkpoint.id,
            checkpoint.import_batch_id,
            checkpoint.entity_id,
            checkpoint.ledger_account_id,
            checkpoint.balance_date,
            checkpoint.balance,
            checkpoint.currency,
            dict(checkpoint.source_metadata or {}),
        )
        for checkpoint in session.scalars(
            select(BalanceCheckpoint)
            .where(BalanceCheckpoint.import_batch_id == batch_id)
            .order_by(BalanceCheckpoint.id)
        ).all()
    ]


def test_parser_reads_signed_decimal_rows_with_data_only(monkeypatch):
    calls = []
    original_load_workbook = openpyxl.load_workbook

    def load_workbook(*args, **kwargs):
        calls.append(kwargs.copy())
        return original_load_workbook(*args, **kwargs)

    monkeypatch.setattr("app.ingestion.baseline.openpyxl.load_workbook", load_workbook)

    records = parse_baseline_xlsx(_valid_bytes())

    assert calls[0]["data_only"] is True
    assert [record.ledger_account_code for record in records] == ["1000", "2000", "3000"]
    assert [record.balance for record in records] == [
        Decimal("100.00"), Decimal("-60.00"), Decimal("-40.00")
    ]
    assert all(record.balance_date == BALANCE_DATE for record in records)
    assert all(record.currency == "INR" for record in records)


def test_valid_baseline_persists_checkpoints_reports_and_activates(baseline_session):
    session, entity_id = baseline_session
    contents = _valid_bytes()
    batch = _stage(session, entity_id, contents)

    report = normalize_balance_checkpoint_xlsx(
        contents,
        entity_id,
        session,
        import_batch_id=batch.id,
        expected_currency="INR",
    )
    session.commit()

    checkpoints = session.scalars(
        select(BalanceCheckpoint).order_by(BalanceCheckpoint.ledger_account_id)
    ).all()
    assert batch.status == "ACTIVE"
    assert batch.raw_source_bytes == contents
    assert [checkpoint.balance for checkpoint in checkpoints] == [
        Decimal("100.00"), Decimal("-60.00"), Decimal("-40.00")
    ]
    assert all(checkpoint.import_batch_id == batch.id for checkpoint in checkpoints)
    assert report["readiness"] == "READY"
    assert report["signed_total"] == "0.00"
    assert report["baseline_coverage"]["complete"] is True
    json.dumps(report)


def test_caller_expected_account_subset_cannot_weaken_master_completeness(baseline_session):
    session, entity_id = baseline_session
    contents = _xlsx_bytes([["1000", BALANCE_DATE, "0.00", "INR"]])
    batch = _stage(session, entity_id, contents)

    with pytest.raises(ValueError, match="Missing expected account code"):
        normalize_balance_checkpoint_xlsx(
            contents,
            entity_id,
            session,
            import_batch_id=batch.id,
            expected_account_codes=["1000"],
            expected_currency="INR",
        )

    assert batch.status == "FAILED"
    assert batch.raw_source_bytes == contents
    assert session.scalar(select(BalanceCheckpoint.id)) is None
    assert batch.validation_report["readiness"] == "INVALID"
    assert batch.validation_report["baseline_coverage"]["complete"] is False
    assert batch.validation_report["missing_account_codes"] == ["2000", "3000"]


def test_tally_balance_checkpoint_is_rejected_without_mutating_staged_batch(baseline_session):
    session, entity_id = baseline_session
    contents = b"not-an-xlsx"
    batch = stage_import_batch(
        session,
        entity_id=entity_id,
        period_start=date(2025, 4, 1),
        period_end=date(2026, 3, 31),
        source="tally",
        source_family="tally",
        original_filename="tally-checkpoint.xlsx",
        contents=contents,
        uploaded_by_user_id=None,
        kind="balance_checkpoint",
        coverage_start=date(2025, 4, 1),
        coverage_end=date(2026, 3, 31),
    )
    original_metadata = dict(batch.source_metadata or {})

    with pytest.raises(ValueError, match="gl_upload"):
        normalize_balance_checkpoint_xlsx(
            contents,
            entity_id,
            session,
            import_batch_id=batch.id,
        )

    assert batch.status == "STAGED"
    assert batch.raw_source_bytes == contents
    assert batch.source_metadata == original_metadata
    assert batch.validation_report == {}
    assert session.scalar(select(BalanceCheckpoint.id)) is None


def test_nonblank_extra_baseline_header_is_rejected():
    contents = _xlsx_bytes(
        [["1000", BALANCE_DATE, "0.00", "INR"]],
        headers=HEADERS + ["Notes"],
    )

    with pytest.raises(ValueError, match="unexpected.*header"):
        parse_baseline_xlsx(contents)


def test_blank_trailing_baseline_headers_are_allowed():
    contents = _xlsx_bytes(
        [["1000", BALANCE_DATE, "0.00", "INR"]],
        headers=HEADERS + [None, None],
    )

    records = parse_baseline_xlsx(contents)

    assert [record.ledger_account_code for record in records] == ["1000"]


def test_duplicate_account_codes_are_rejected():
    contents = _xlsx_bytes([
        ["1000", BALANCE_DATE, "50.00", "INR"],
        ["1000", BALANCE_DATE, "-50.00", "INR"],
    ])

    with pytest.raises(ValueError, match="[Dd]uplicate account code"):
        parse_baseline_xlsx(contents)


@pytest.mark.parametrize(
    ("header_index", "unsupported_header"),
    [
        (0, "G/L Account"),
        (0, "G/L Account Code"),
        (2, "Closing Balance"),
        (2, "Balance"),
        (3, "Currency Code"),
    ],
)
def test_unsupported_baseline_header_aliases_are_rejected(header_index, unsupported_header):
    headers = HEADERS.copy()
    headers[header_index] = unsupported_header
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(headers)
    worksheet.append(["1000", BALANCE_DATE, "100.00", "INR"])
    output = io.BytesIO()
    workbook.save(output)

    with pytest.raises(ValueError, match="exact required headers"):
        parse_baseline_xlsx(output.getvalue())


@pytest.mark.parametrize("header_row", [2, 16])
def test_baseline_headers_outside_fixed_row_one_are_rejected(header_row):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    for _ in range(header_row - 1):
        worksheet.append(["Report title"])
    worksheet.append(HEADERS)
    worksheet.append(["1000", BALANCE_DATE, "100.00", "INR"])
    output = io.BytesIO()
    workbook.save(output)

    with pytest.raises(ValueError, match="row 1"):
        parse_baseline_xlsx(output.getvalue())


def test_non_balanced_signed_total_is_rejected():
    contents = _xlsx_bytes([
        ["1000", BALANCE_DATE, "100.00", "INR"],
        ["2000", BALANCE_DATE, "-99.98", "INR"],
    ])

    with pytest.raises(ValueError, match="signed total"):
        parse_baseline_xlsx(contents)


def test_date_and_currency_inconsistency_are_rejected():
    with pytest.raises(ValueError, match="one balance date"):
        parse_baseline_xlsx(_xlsx_bytes([
            ["1000", BALANCE_DATE, "100.00", "INR"],
            ["2000", date(2025, 3, 30), "-100.00", "INR"],
        ]))

    with pytest.raises(ValueError, match="same currency"):
        parse_baseline_xlsx(_xlsx_bytes([
            ["1000", BALANCE_DATE, "100.00", "INR"],
            ["2000", BALANCE_DATE, "-100.00", "USD"],
        ]))


def test_unknown_and_missing_master_accounts_are_rejected(baseline_session):
    session, entity_id = baseline_session
    contents = _xlsx_bytes([
        ["1000", BALANCE_DATE, "100.00", "INR"],
        ["9999", BALANCE_DATE, "-100.00", "INR"],
    ])
    batch = _stage(session, entity_id, contents)

    with pytest.raises(ValueError, match="Unknown account code.*9999"):
        normalize_balance_checkpoint_xlsx(
            contents,
            entity_id,
            session,
            import_batch_id=batch.id,
            expected_currency="INR",
        )

    assert batch.status == "FAILED"
    assert session.scalar(select(BalanceCheckpoint.id)) is None
    assert batch.raw_source_bytes == contents
    errors = " ".join(batch.validation_report["errors"])
    assert "Missing expected account code" in errors
    assert "2000" in errors and "3000" in errors


def test_pdf_only_input_is_rejected_but_retained_as_evidence(baseline_session):
    session, entity_id = baseline_session
    pdf_bytes = b"%PDF-1.7 signed closing trial balance"
    batch = _stage(session, entity_id, pdf_bytes)

    with pytest.raises(ValueError, match="PDF.*structured XLSX"):
        normalize_balance_checkpoint_xlsx(
            pdf_bytes,
            entity_id,
            session,
            import_batch_id=batch.id,
        )

    assert batch.status == "FAILED"
    assert batch.raw_source_bytes == pdf_bytes
    evidence = batch.source_metadata["supporting_evidence"]["signed_pdf"]
    assert base64.b64decode(evidence["bytes_base64"]) == pdf_bytes
    assert session.scalar(select(BalanceCheckpoint.id)) is None


def test_invalid_staged_batch_keeps_workbook_and_optional_pdf_bytes(baseline_session):
    session, entity_id = baseline_session
    contents = _xlsx_bytes([
        ["1000", BALANCE_DATE, "10.00", "INR"],
        ["2000", BALANCE_DATE, "-9.00", "INR"],
    ])
    pdf_bytes = b"%PDF-1.7 signed evidence"
    batch = _stage(session, entity_id, contents)

    with pytest.raises(ValueError, match="signed total"):
        normalize_balance_checkpoint_xlsx(
            contents,
            entity_id,
            session,
            import_batch_id=batch.id,
            signed_pdf_bytes=pdf_bytes,
            expected_currency="INR",
        )

    assert batch.status == "FAILED"
    assert batch.raw_source_bytes == contents
    evidence = batch.source_metadata["supporting_evidence"]["signed_pdf"]
    assert base64.b64decode(evidence["bytes_base64"]) == pdf_bytes


def test_mismatched_caller_bytes_fail_staged_batch_and_retain_artifacts(baseline_session):
    session, entity_id = baseline_session
    staged_contents = _valid_bytes()
    caller_contents = _xlsx_bytes([
        ["1000", BALANCE_DATE, "90.00", "INR"],
        ["2000", BALANCE_DATE, "-50.00", "INR"],
        ["3000", BALANCE_DATE, "-40.00", "INR"],
    ])
    pdf_bytes = b"%PDF-1.7 signed evidence"
    batch = _stage(session, entity_id, staged_contents)

    with pytest.raises(ValueError, match="do not match staged"):
        normalize_balance_checkpoint_xlsx(
            caller_contents,
            entity_id,
            session,
            import_batch_id=batch.id,
            signed_pdf_bytes=pdf_bytes,
            expected_currency="INR",
        )

    assert batch.status == "FAILED"
    assert batch.raw_source_bytes == staged_contents
    evidence = batch.source_metadata["supporting_evidence"]["signed_pdf"]
    assert base64.b64decode(evidence["bytes_base64"]) == pdf_bytes
    assert session.scalar(select(BalanceCheckpoint.id)) is None
    assert any("do not match staged" in error for error in batch.validation_report["errors"])


@pytest.mark.parametrize(
    ("rows", "error_match"),
    [
        (
            [
                ["1000", BALANCE_DATE, "100.00", "INR"],
                [None, BALANCE_DATE, "-60.00", "INR"],
                ["2000", BALANCE_DATE, "-40.00", "INR"],
            ],
            "Account Code is a required value",
        ),
        (
            [
                ["1000", BALANCE_DATE, "100.00", "INR"],
                ["2000", BALANCE_DATE, "-50.00", "INR"],
                ["3000", BALANCE_DATE, "-40.00", "INR"],
            ],
            "signed total",
        ),
        (
            [
                ["1000", BALANCE_DATE, "100.00", "INR"],
                ["9999", BALANCE_DATE, "-100.00", "INR"],
            ],
            "Unknown account code",
        ),
    ],
)
def test_baseline_failure_report_accounts_every_source_row(baseline_session, rows, error_match):
    session, entity_id = baseline_session
    contents = _xlsx_bytes(rows)
    batch = _stage(session, entity_id, contents)

    with pytest.raises(ValueError, match=error_match):
        normalize_balance_checkpoint_xlsx(
            contents,
            entity_id,
            session,
            import_batch_id=batch.id,
            expected_currency="INR",
        )

    report = batch.validation_report
    assert report["input_rows"] == len(rows)
    assert report["accepted_rows"] + report["skipped_rows"] + report["rejected_rows"] == report["input_rows"]
    assert report["accepted_rows"] == 0
    assert report["rejected_rows"] == len(rows)
    assert {reason["row"] for reason in report["reject_reasons"]} == set(range(2, len(rows) + 2))
    assert all(reason["reason"] for reason in report["reject_reasons"])


def test_active_exact_sha_duplicate_replay_is_idempotent(baseline_session):
    session, entity_id = baseline_session
    contents = _valid_bytes()
    batch = _stage(session, entity_id, contents)
    first_report = normalize_balance_checkpoint_xlsx(
        contents, entity_id, session, import_batch_id=batch.id, expected_currency="INR"
    )
    session.commit()
    checkpoint_ids = [checkpoint.id for checkpoint in session.scalars(select(BalanceCheckpoint)).all()]

    duplicate = _stage(session, entity_id, contents)
    duplicate_id = duplicate.id
    session.expunge_all()
    replay_report = normalize_balance_checkpoint_xlsx(
        contents, entity_id, session, import_batch_id=duplicate_id, expected_currency="INR"
    )
    session.commit()

    replayed_batch = session.get(ImportBatch, duplicate_id)
    assert replayed_batch.id == batch.id
    assert replayed_batch.status == "ACTIVE"
    assert replay_report == first_report
    assert [checkpoint.id for checkpoint in session.scalars(select(BalanceCheckpoint)).all()] == checkpoint_ids


@pytest.mark.parametrize("mutation", ["missing", "altered"])
def test_active_baseline_replay_rejects_corrupt_children_without_repopulation(
    baseline_session, mutation
):
    session, entity_id = baseline_session
    contents = _valid_bytes()
    batch = _stage(session, entity_id, contents)
    first_report = normalize_balance_checkpoint_xlsx(
        contents, entity_id, session, import_batch_id=batch.id, expected_currency="INR"
    )
    session.commit()

    checkpoint = session.scalars(
        select(BalanceCheckpoint).where(BalanceCheckpoint.import_batch_id == batch.id)
    ).first()
    assert checkpoint is not None
    if mutation == "missing":
        session.delete(checkpoint)
    else:
        checkpoint.balance = Decimal("999.00")
    session.commit()
    corrupted_count = session.scalar(
        select(func.count()).select_from(BalanceCheckpoint).where(
            BalanceCheckpoint.import_batch_id == batch.id
        )
    )
    corrupted_snapshot = _checkpoint_snapshot(session, batch.id)

    duplicate = _stage(session, entity_id, contents)
    assert duplicate.id == batch.id
    with pytest.raises(ValueError, match="missing or altered"):
        normalize_balance_checkpoint_xlsx(
            contents, entity_id, session, import_batch_id=duplicate.id, expected_currency="INR"
        )

    session.commit()
    assert session.get(ImportBatch, batch.id).status == "ACTIVE"
    assert session.get(ImportBatch, batch.id).validation_report == first_report
    assert session.scalar(
        select(func.count()).select_from(BalanceCheckpoint).where(
            BalanceCheckpoint.import_batch_id == batch.id
        )
    ) == corrupted_count
    assert _checkpoint_snapshot(session, batch.id) == corrupted_snapshot


def test_active_baseline_replay_rejects_foreign_entity_checkpoint_without_mutation(baseline_session):
    session, entity_id = baseline_session
    contents = _valid_bytes()
    batch = _stage(session, entity_id, contents)
    first_report = normalize_balance_checkpoint_xlsx(
        contents, entity_id, session, import_batch_id=batch.id, expected_currency="INR"
    )
    session.commit()

    organization_id = session.get(Entity, entity_id).organization_id
    foreign_entity = Entity(
        organization_id=organization_id,
        name="Foreign Baseline Owner",
        materiality_threshold=Decimal("0.00"),
    )
    session.add(foreign_entity)
    session.flush()
    checkpoint = session.scalars(
        select(BalanceCheckpoint).where(BalanceCheckpoint.import_batch_id == batch.id)
    ).first()
    assert checkpoint is not None
    checkpoint.entity_id = foreign_entity.id
    session.commit()
    corrupted_snapshot = _checkpoint_snapshot(session, batch.id)

    duplicate = _stage(session, entity_id, contents)
    assert duplicate.id == batch.id
    with pytest.raises(ValueError, match="missing or altered"):
        normalize_balance_checkpoint_xlsx(
            contents, entity_id, session, import_batch_id=duplicate.id, expected_currency="INR"
        )

    session.commit()
    persisted = session.get(ImportBatch, batch.id)
    assert persisted.status == "ACTIVE"
    assert persisted.validation_report == first_report
    assert _checkpoint_snapshot(session, batch.id) == corrupted_snapshot


def test_active_exact_sha_journal_batch_is_rejected_by_baseline_normalizer(baseline_session):
    session, entity_id = baseline_session
    contents = _valid_bytes()
    batch = stage_import_batch(
        session,
        entity_id=entity_id,
        period_start=date(2025, 4, 1),
        period_end=date(2026, 3, 31),
        source="gl_upload",
        source_family="gl_upload",
        original_filename="journal.xlsx",
        contents=contents,
        uploaded_by_user_id=None,
        kind="journal",
        coverage_start=date(2025, 4, 1),
        coverage_end=date(2026, 3, 31),
    )
    activate_import_batch(session, batch, validation_report={"readiness": "READY"})

    with pytest.raises(ValueError, match="balance_checkpoint"):
        normalize_balance_checkpoint_xlsx(
            contents, entity_id, session, import_batch_id=batch.id, expected_currency="INR"
        )

    assert batch.status == "ACTIVE"
    assert session.scalar(select(BalanceCheckpoint.id)) is None


def test_baseline_fingerprint_includes_account_balance_pairs(baseline_session):
    session, entity_id = baseline_session
    first_contents = _valid_bytes()
    second_contents = _xlsx_bytes([
        ["1000", BALANCE_DATE, "90.00", "INR"],
        ["2000", BALANCE_DATE, "-50.00", "INR"],
        ["3000", BALANCE_DATE, "-40.00", "INR"],
    ])
    first = _stage(session, entity_id, first_contents)
    second = _stage(session, entity_id, second_contents)

    first_report = normalize_balance_checkpoint_xlsx(
        first_contents, entity_id, session, import_batch_id=first.id, expected_currency="INR"
    )
    second_report = normalize_balance_checkpoint_xlsx(
        second_contents, entity_id, session, import_batch_id=second.id, expected_currency="INR"
    )

    assert first_report["signed_total"] == second_report["signed_total"] == "0.00"
    assert first_report["dataset_fingerprint"] != second_report["dataset_fingerprint"]
