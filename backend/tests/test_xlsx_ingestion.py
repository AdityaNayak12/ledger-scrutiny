import os
import io
from datetime import date
from decimal import Decimal
import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from app.main import app
from app.db.base import Base
from app.db.models import Entity, ImportBatch, JournalEntry, JournalLine, LedgerAccount, Organization
from app.ingestion.batches import stage_import_batch
from app.ingestion.xlsx_normalizer import normalize_gl_xlsx, normalize_xlsx_confirm, parse_gl_xlsx
from conftest import TestingSessionLocal

client = TestClient(app)

GOLDEN_GL = os.path.join(os.path.dirname(__file__), "fixtures/golden_gl.xlsx")


def _xlsx_bytes(headers, rows, preamble=()):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    for row in preamble:
        worksheet.append(row)
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.fixture
def canonical_session():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        organization = Organization(name="GL Test Org")
        session.add(organization)
        session.flush()
        entity = Entity(
            organization_id=organization.id,
            name="GL Test Entity",
            materiality_threshold=Decimal("0.00"),
        )
        session.add(entity)
        session.flush()
        yield session, entity.id
    Base.metadata.drop_all(engine)


def _stage_gl_batch(session, entity_id, contents):
    return stage_import_batch(
        session,
        entity_id=entity_id,
        period_start=date(2025, 4, 1),
        period_end=date(2026, 3, 31),
        source="gl_upload",
        source_family="gl_upload",
        original_filename="golden_gl.xlsx",
        contents=contents,
        uploaded_by_user_id=None,
        coverage_start=date(2025, 4, 1),
        coverage_end=date(2026, 3, 31),
    )


def test_fixed_gl_normalizer_persists_golden_canonical_counts_and_report(canonical_session):
    session, entity_id = canonical_session
    contents = open(GOLDEN_GL, "rb").read()
    batch = _stage_gl_batch(session, entity_id, contents)

    report = normalize_gl_xlsx(
        contents,
        target_period_start=date(2025, 4, 1),
        target_period_end=date(2026, 3, 31),
        entity_id=entity_id,
        session=session,
        import_batch_id=batch.id,
    )
    session.commit()

    assert session.scalar(select(func.count()).select_from(JournalLine)) == 6
    assert session.scalar(select(func.count()).select_from(JournalEntry)) == 3
    assert session.scalar(select(func.count()).select_from(LedgerAccount)) == 5
    assert [line.side for line in session.scalars(select(JournalLine).order_by(JournalLine.source_row_number))] == [
        "debit", "credit", "debit", "credit", "debit", "credit"
    ]
    assert report["accepted_rows"] == 6
    assert report["document_count"] == 3
    assert report["unbalanced_document_count"] == 0
    assert Decimal(report["net"]) == Decimal("0")
    assert batch.validation_report["readiness"] == "PARTIAL"
    assert batch.validation_report["account_codes"] == ["110000", "120000", "210000", "400000", "500000"]


def test_fixed_gl_parser_uses_data_only_and_scans_only_first_fifteen_rows(monkeypatch):
    headers = ["Document Number", "G/L Account", "Posting Date", "Amount in local currency"]
    contents = _xlsx_bytes(
        headers,
        [
            ["DOC-1", "1000", date(2025, 4, 1), 1],
            ["DOC-1", "2000", date(2025, 4, 1), -1],
        ],
        preamble=[["Report title"] for _ in range(14)],
    )
    import app.ingestion.xlsx_normalizer as normalizer

    calls = []
    original_load_workbook = normalizer.openpyxl.load_workbook

    def load_workbook(*args, **kwargs):
        calls.append(kwargs.copy())
        return original_load_workbook(*args, **kwargs)

    monkeypatch.setattr(normalizer.openpyxl, "load_workbook", load_workbook)
    entries = parse_gl_xlsx(contents, date(2025, 4, 1), date(2025, 4, 30))

    assert len(entries) == 1
    assert calls[0]["data_only"] is True

    after_fifteen_rows = _xlsx_bytes(
        headers,
        [
            ["DOC-1", "1000", date(2025, 4, 1), 1],
            ["DOC-1", "2000", date(2025, 4, 1), -1],
        ],
        preamble=[["Report title"] for _ in range(15)],
    )
    with pytest.raises(ValueError, match="first 15 rows"):
        parse_gl_xlsx(after_fifteen_rows, date(2025, 4, 1), date(2025, 4, 30))

    wrong_header = _xlsx_bytes(
        ["Document Number", "G/L account", "Posting Date", "Amount in local currency"],
        [["DOC-1", "1000", date(2025, 4, 1), 1], ["DOC-1", "2000", date(2025, 4, 1), -1]],
    )
    with pytest.raises(ValueError, match="exact required GL headers"):
        parse_gl_xlsx(wrong_header, date(2025, 4, 1), date(2025, 4, 30))


def test_fixed_gl_parser_never_evaluates_formula_values():
    contents = _xlsx_bytes(
        ["Document Number", "G/L Account", "Posting Date", "Amount in local currency"],
        [["DOC-1", "1000", date(2025, 4, 1), "=1+1"]],
    )

    with pytest.raises(ValueError, match="Amount in local currency.*required"):
        parse_gl_xlsx(contents, date(2025, 4, 1), date(2025, 4, 30))


def test_fixed_gl_parser_preserves_dimensions_and_source_row(canonical_session):
    session, entity_id = canonical_session
    headers = [
        "Document Number", "G/L Account", "Posting Date", "Amount in local currency",
        "Document Date", "Document Type", "Posting Key", "Invoice Reference",
        "Clearing Document", "Profit Centre", "Cost Centre", "Text", "Supplier",
        "WBS Element", "Purchasing Document", "Customer", "Business Area",
    ]
    contents = _xlsx_bytes(headers, [
        [
            "DOC-1", "1000", date(2025, 4, 1), 100,
            date(2025, 4, 1), "SA", "40", "INV-1", "CLR-1", "PC-1", "CC-1",
            "Narration", "V-1", "WBS-1", "PO-1", "C-1", "BA-1",
        ],
        [
            "DOC-1", "2000", date(2025, 4, 1), -100,
            date(2025, 4, 1), "SA", "50", "INV-1", "CLR-1", "PC-1", "CC-1",
            "Narration", "V-1", "WBS-1", "PO-1", "C-1", "BA-1",
        ],
    ])
    entries = parse_gl_xlsx(contents, date(2025, 4, 1), date(2025, 4, 30))
    entry = entries[0]
    line = entry.lines[0]

    assert line.source_row_number == 2
    assert entry.document_date == date(2025, 4, 1)
    assert entry.document_type == "SA"
    assert entry.narration == "Narration"
    assert line.posting_key == "40"
    assert line.reference == "INV-1"
    assert line.clearing_document == "CLR-1"
    assert line.profit_center == "PC-1"
    assert line.cost_center == "CC-1"
    assert line.text == "Narration"
    assert line.supplier == "V-1"
    assert line.wbs == "WBS-1"
    assert line.purchasing_document == "PO-1"
    assert line.customer == "C-1"
    assert line.dimensions["Business Area"] == "BA-1"
    assert line.source_metadata["Business Area"] == "BA-1"

    batch = _stage_gl_batch(session, entity_id, contents)
    normalize_gl_xlsx(
        contents,
        date(2025, 4, 1),
        date(2025, 4, 30),
        entity_id,
        session,
        import_batch_id=batch.id,
    )
    persisted = session.scalar(select(JournalLine))
    assert persisted.source_row_number == 2
    assert persisted.dimensions["Business Area"] == "BA-1"
    assert persisted.source_metadata["G/L Account"] == "1000"


@pytest.mark.parametrize(
    ("document", "account", "posting_date", "amount", "message"),
    [
        (None, "1000", date(2025, 4, 1), 100, "Document Number.*required"),
        ("DOC-1", None, date(2025, 4, 1), 100, "G/L Account.*required"),
        ("DOC-1", "1000", "not-a-date", 100, "Posting Date.*date"),
        ("DOC-1", "1000", date(2025, 4, 1), "not-a-decimal", "Amount in local currency.*decimal"),
    ],
)
def test_fixed_gl_parser_rejects_malformed_required_values(
    document, account, posting_date, amount, message
):
    contents = _xlsx_bytes(
        ["Document Number", "G/L Account", "Posting Date", "Amount in local currency"],
        [[document, account, posting_date, amount]],
    )

    with pytest.raises(ValueError, match=message):
        parse_gl_xlsx(contents, date(2025, 4, 1), date(2025, 4, 30))


def test_fixed_gl_parser_rejects_unbalanced_documents():
    contents = _xlsx_bytes(
        ["Document Number", "G/L Account", "Posting Date", "Amount in local currency"],
        [
            ["DOC-1", "1000", date(2025, 4, 1), 100],
            ["DOC-1", "2000", date(2025, 4, 1), -99.98],
        ],
    )

    with pytest.raises(ValueError, match="DOC-1.*balance"):
        parse_gl_xlsx(contents, date(2025, 4, 1), date(2025, 4, 30))


def test_fixed_gl_parser_rejects_rows_outside_declared_coverage():
    contents = _xlsx_bytes(
        ["Document Number", "G/L Account", "Posting Date", "Amount in local currency"],
        [["DOC-1", "1000", date(2025, 5, 1), 100]],
    )

    with pytest.raises(ValueError, match="outside declared coverage"):
        parse_gl_xlsx(contents, date(2025, 4, 1), date(2025, 4, 30))


def test_fixed_gl_normalizer_reports_skips_formula_cache_warnings_and_unclassified_accounts(canonical_session):
    session, entity_id = canonical_session
    headers = [
        "Document Number", "G/L Account", "Posting Date", "Amount in local currency", "Text",
    ]
    contents = _xlsx_bytes(headers, [
        ["DOC-1", "1000", date(2025, 4, 1), 100, '=CONCAT("not", "cached")'],
        [None, None, None, None, None],
        ["DOC-1", "2000", date(2025, 4, 1), -100, '=CONCAT("not", "cached")'],
    ])
    batch = _stage_gl_batch(session, entity_id, contents)

    report = normalize_gl_xlsx(
        contents,
        date(2025, 4, 1),
        date(2025, 4, 30),
        entity_id,
        session,
        import_batch_id=batch.id,
    )

    assert report["accepted_rows"] == 2
    assert report["skipped_rows"] == 1
    assert report["skip_reasons"] == [{"row": 3, "reason": "blank source row"}]
    assert any("missing cached formula" in warning for warning in report["warnings"])
    assert report["unmapped_account_codes"] == ["1000", "2000"]
    assert all(account.group_name is None for account in session.scalars(select(LedgerAccount)))


def test_fixed_gl_parser_allows_repeated_account_codes_within_a_document():
    contents = _xlsx_bytes(
        ["Document Number", "G/L Account", "Posting Date", "Amount in local currency"],
        [
            ["DOC-1", "1000", date(2025, 4, 1), 60],
            ["DOC-1", "1000", date(2025, 4, 1), 40],
            ["DOC-1", "2000", date(2025, 4, 1), -100],
        ],
    )

    entries = parse_gl_xlsx(contents, date(2025, 4, 1), date(2025, 4, 30))

    assert len(entries) == 1
    assert [line.ledger_account_code for line in entries[0].lines] == ["1000", "1000", "2000"]


@pytest.mark.skipif(
    not os.path.exists("/Users/adinayak18/Downloads/GL Dump Q1.XLSX"),
    reason="supplied Q1 workbook is not available",
)
def test_supplied_q1_gl_workbook_is_accepted_with_footer_and_nonnumeric_quantity():
    contents = open("/Users/adinayak18/Downloads/GL Dump Q1.XLSX", "rb").read()

    entries = parse_gl_xlsx(contents, date(2025, 4, 1), date(2025, 6, 30))

    assert sum(len(entry.lines) for entry in entries) == 59167
    assert len(entries) == 10914
    assert len({line.ledger_account_code for entry in entries for line in entry.lines}) == 445
    assert sum((line.amount for entry in entries for line in entry.lines), Decimal("0")) == Decimal("0.00")
    om_line = next(
        line for entry in entries for line in entry.lines
        if line.source_row_number == 28096
    )
    assert om_line.quantity is None
    assert om_line.source_metadata["Quantity"] == "OM"
    assert om_line.dimensions["Quantity"] == "OM"


def test_fixed_gl_parser_preserves_nonnumeric_optional_quantity_and_warns(canonical_session):
    session, entity_id = canonical_session
    contents = _xlsx_bytes(
        ["Document Number", "G/L Account", "Posting Date", "Amount in local currency", "Quantity"],
        [["DOC-1", "1000", date(2025, 4, 1), 1, "OM"],
         ["DOC-1", "2000", date(2025, 4, 1), -1, 2]],
    )
    batch = _stage_gl_batch(session, entity_id, contents)

    report = normalize_gl_xlsx(
        contents, date(2025, 4, 1), date(2025, 4, 30), entity_id, session,
        import_batch_id=batch.id,
    )

    line = session.scalar(select(JournalLine).where(JournalLine.source_row_number == 2))
    assert line.quantity is None
    assert line.source_metadata["Quantity"] == "OM"
    assert line.dimensions["Quantity"] == "OM"
    assert sum("non-numeric" in warning for warning in report["warnings"]) == 1


def test_fixed_gl_parser_skips_only_structured_summary_rows(canonical_session):
    session, entity_id = canonical_session
    headers = [
        "Document Number", "G/L Account", "Posting Date", "Amount in local currency", "Description",
    ]
    contents = _xlsx_bytes(headers, [
        ["DOC-1", "1000", date(2025, 4, 1), 1, "valid"],
        ["DOC-1", "2000", date(2025, 4, 1), -1, "valid"],
        [None, None, None, 0, "LIABILITY TOTAL"],
    ])

    parsed = parse_gl_xlsx(contents, date(2025, 4, 1), date(2025, 4, 30))
    assert len(parsed) == 1
    batch = _stage_gl_batch(session, entity_id, contents)
    report = normalize_gl_xlsx(
        contents, date(2025, 4, 1), date(2025, 4, 30), entity_id, session,
        import_batch_id=batch.id,
    )
    assert report["skipped_rows"] == 1
    assert report["skip_reasons"] == [
        {"row": 4, "reason": "summary/footer row", "column": "Description", "marker": "LIABILITY TOTAL"}
    ]

    malformed = _xlsx_bytes(headers, [[None, None, None, 1, "ordinary text"]])
    with pytest.raises(ValueError, match="Document Number.*required"):
        parse_gl_xlsx(malformed, date(2025, 4, 1), date(2025, 4, 30))

    blank_amount_footer = _xlsx_bytes(
        headers, [["DOC-2", "1000", date(2025, 4, 1), None, "LIABILITY TOTAL"]]
    )
    with pytest.raises(ValueError, match="Amount in local currency.*required"):
        parse_gl_xlsx(blank_amount_footer, date(2025, 4, 1), date(2025, 4, 30))


def test_fixed_gl_normalizer_resolves_accounts_by_entity_scoped_external_code_only(canonical_session):
    session, entity_id = canonical_session
    existing = LedgerAccount(
        entity_id=entity_id,
        external_code="LEGACY-1000",
        name="1000",
        group_name="Current Assets",
        normal_balance="debit",
    )
    session.add(existing)
    session.flush()
    contents = _xlsx_bytes(
        ["Document Number", "G/L Account", "Posting Date", "Amount in local currency"],
        [["DOC-1", "1000", date(2025, 4, 1), 1], ["DOC-1", "2000", date(2025, 4, 1), -1]],
    )
    batch = _stage_gl_batch(session, entity_id, contents)

    normalize_gl_xlsx(
        contents, date(2025, 4, 1), date(2025, 4, 30), entity_id, session,
        import_batch_id=batch.id,
    )

    accounts = session.scalars(select(LedgerAccount).order_by(LedgerAccount.id)).all()
    assert existing.external_code == "LEGACY-1000"
    account_1000 = next(account for account in accounts if account.external_code == "1000")
    assert account_1000.id != existing.id


@pytest.mark.parametrize("status", ["ACTIVE", "FAILED", "SUPERSEDED"])
def test_fixed_gl_normalizer_requires_staged_batch_and_preserves_existing_report(canonical_session, status):
    session, entity_id = canonical_session
    contents = _xlsx_bytes(
        ["Document Number", "G/L Account", "Posting Date", "Amount in local currency"],
        [["DOC-1", "1000", date(2025, 4, 1), 1], ["DOC-1", "2000", date(2025, 4, 1), -1]],
    )
    batch = _stage_gl_batch(session, entity_id, contents)
    batch.status = status
    batch.validation_report = {"sentinel": status}
    session.flush()

    with pytest.raises(ValueError, match="requires a STAGED ImportBatch"):
        normalize_gl_xlsx(
            contents, date(2025, 4, 1), date(2025, 4, 30), entity_id, session,
            import_batch_id=batch.id,
        )

    assert batch.validation_report == {"sentinel": status}
    assert session.scalar(select(func.count()).select_from(JournalLine)) == 0
    assert session.scalar(select(func.count()).select_from(LedgerAccount)) == 0


def test_fixed_gl_normalizer_records_structured_report_for_staged_failure(canonical_session):
    session, entity_id = canonical_session
    contents = _xlsx_bytes(
        ["Document Number", "G/L Account", "Posting Date", "Amount in local currency", "Description"],
        [[None, None, None, 0, "LIABILITY TOTAL"], ["DOC-1", None, date(2025, 4, 1), 1, "bad"]],
    )
    batch = _stage_gl_batch(session, entity_id, contents)

    with pytest.raises(ValueError, match="G/L Account.*required"):
        normalize_gl_xlsx(
            contents, date(2025, 4, 1), date(2025, 4, 30), entity_id, session,
            import_batch_id=batch.id,
        )

    assert batch.status == "FAILED"
    assert batch.validation_report["input_rows"] == 2
    assert batch.validation_report["accepted_rows"] == 0
    assert batch.validation_report["skipped_rows"] == 1
    assert batch.validation_report["rejected_rows"] == 1
    assert batch.validation_report["skip_reasons"][0]["row"] == 2
    assert batch.validation_report["reject_reasons"][0]["row"] == 3


@pytest.mark.parametrize("sign_convention", ["positive_is_credit", "separate_dr_cr_columns"])
def test_fixed_gl_confirm_rejects_unsupported_sign_convention(canonical_session, sign_convention):
    session, entity_id = canonical_session
    contents = _xlsx_bytes(
        ["Document Number", "G/L Account", "Posting Date", "Amount in local currency"],
        [["DOC-1", "1000", date(2025, 4, 1), 100], ["DOC-1", "2000", date(2025, 4, 1), -100]],
    )
    batch = _stage_gl_batch(session, entity_id, contents)

    with pytest.raises(ValueError, match="signed amounts"):
        normalize_xlsx_confirm(
            contents,
            {},
            sign_convention,
            date(2025, 4, 1),
            date(2026, 3, 31),
            entity_id,
            session,
            import_batch_id=batch.id,
        )

    assert session.scalar(select(func.count()).select_from(JournalLine)) == 0
    assert session.scalar(select(func.count()).select_from(JournalEntry)) == 0


def test_fixed_gl_parser_aggregates_uncached_formula_warnings_for_unknown_columns(canonical_session):
    session, entity_id = canonical_session
    contents = _xlsx_bytes(
        ["Document Number", "G/L Account", "Posting Date", "Amount in local currency", "Mystery Dimension"],
        [["DOC-1", "1000", date(2025, 4, 1), 1, "=CONCAT(\"a\", \"b\")"],
         ["DOC-1", "2000", date(2025, 4, 1), -1, "=CONCAT(\"c\", \"d\")"]],
    )

    batch = _stage_gl_batch(session, entity_id, contents)
    report = normalize_gl_xlsx(
        contents, date(2025, 4, 1), date(2025, 4, 30), entity_id, session,
        import_batch_id=batch.id,
    )

    assert len([warning for warning in report["warnings"] if "missing cached formula" in warning]) == 1
    assert "Mystery Dimension" in next(
        warning for warning in report["warnings"] if "missing cached formula" in warning
    )


def test_fixed_gl_parser_rejects_amount_precision_that_numeric_20_2_cannot_retain():
    contents = _xlsx_bytes(
        ["Document Number", "G/L Account", "Posting Date", "Amount in local currency"],
        [["DOC-1", "1000", date(2025, 4, 1), "1.001"]],
    )

    with pytest.raises(ValueError, match="Amount in local currency.*2 decimal"):
        parse_gl_xlsx(contents, date(2025, 4, 1), date(2025, 4, 30))


def test_fixed_gl_parser_rejects_numeric_quantity_precision_beyond_numeric_20_4():
    contents = _xlsx_bytes(
        ["Document Number", "G/L Account", "Posting Date", "Amount in local currency", "Quantity"],
        [["DOC-1", "1000", date(2025, 4, 1), 1, "1.00001"],
         ["DOC-1", "2000", date(2025, 4, 1), -1, 0]],
    )

    with pytest.raises(ValueError, match="Quantity.*4 decimal"):
        parse_gl_xlsx(contents, date(2025, 4, 1), date(2025, 4, 30))


@pytest.fixture
def auth_headers():
    res = client.post("/auth/register", json={
        "organization_name": "XLSX Test Firm",
        "email": f"xlsx.auditor.{os.urandom(4).hex()}@firm.com",
        "password": "Password123!"
    })
    assert res.status_code == 201
    token = res.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def test_entity(auth_headers):
    res = client.post("/entities", json={
        "name": "XLSX Test Entity",
        "materiality_threshold": 0.0
    }, headers=auth_headers)
    assert res.status_code == 201
    return res.json()["id"]


def test_fixed_gl_confirm_activates_after_normalization_and_rolls_back_invalid_replacement(
    auth_headers, test_entity
):
    valid_contents = _xlsx_bytes(
        ["Document Number", "G/L Account", "Posting Date", "Amount in local currency"],
        [["DOC-1", "1000", date(2025, 4, 1), 100],
         ["DOC-1", "2000", date(2025, 4, 1), -100]],
    )
    upload = client.post(
        f"/entities/{test_entity}/upload-xlsx/confirm",
        data={
            "column_mapping": "{}",
            "target_period_start": "2025-04-01",
            "target_period_end": "2026-03-31",
        },
        files={"file": ("fixed_gl.xlsx", io.BytesIO(valid_contents),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=auth_headers,
    )

    assert upload.status_code == 200, upload.text
    response = upload.json()
    assert response["status"] == "ACTIVE"
    assert response["validation_report"]["parser"] == "xlsx_gl"
    assert response["validation_report"]["dataset_fingerprint"]
    batch_id = response["import_batch_id"]
    with TestingSessionLocal() as session:
        assert session.get(ImportBatch, batch_id).status == "ACTIVE"

    duplicate = client.post(
        f"/entities/{test_entity}/upload-xlsx/confirm",
        data={
            "column_mapping": "{}",
            "target_period_start": "2025-04-01",
            "target_period_end": "2026-03-31",
        },
        files={"file": ("fixed_gl.xlsx", io.BytesIO(valid_contents),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=auth_headers,
    )
    assert duplicate.status_code == 200, duplicate.text
    assert duplicate.json()["import_batch_id"] == batch_id

    invalid_contents = _xlsx_bytes(
        ["Document Number", "G/L Account", "Posting Date", "Amount in local currency"],
        [["DOC-2", "1000", date(2025, 4, 1), 10],
         ["DOC-2", "2000", date(2025, 4, 1), -9]],
    )
    failed = client.post(
        f"/entities/{test_entity}/upload-xlsx/confirm",
        data={
            "column_mapping": "{}",
            "target_period_start": "2025-04-01",
            "target_period_end": "2026-03-31",
        },
        files={"file": ("replacement.xlsx", io.BytesIO(invalid_contents),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=auth_headers,
    )
    assert failed.status_code == 400

    with TestingSessionLocal() as session:
        active = session.get(ImportBatch, batch_id)
        assert active is not None, session.scalars(select(ImportBatch)).all()
        assert active.status == "ACTIVE"
        assert active.validation_report["dataset_fingerprint"] == response["validation_report"]["dataset_fingerprint"]
        assert session.scalar(select(func.count()).select_from(JournalLine)) == 2
        assert session.scalar(select(func.count()).select_from(ImportBatch)) == 2
        failed_batch = session.scalar(
            select(ImportBatch).where(
                ImportBatch.entity_id == test_entity,
                ImportBatch.status == "FAILED",
            )
        )
        assert failed_batch is not None
        assert failed_batch.raw_source_bytes == invalid_contents
        assert failed_batch.validation_report["rejected_rows"] == 1
        assert failed_batch.validation_report["errors"]
        assert failed_batch.validation_report["reject_reasons"]


def test_xlsx_confirm_persists_fixed_canonical_fixture(auth_headers, test_entity):
    contents = open(GOLDEN_GL, "rb").read()
    confirm_res = client.post(
        f"/entities/{test_entity}/upload-xlsx/confirm",
        data={
            "column_mapping": "{}",
            "target_period_start": "2025-04-01",
            "target_period_end": "2026-03-31",
            "clear_only_period": "true",
        },
        files={"file": ("golden_gl.xlsx", io.BytesIO(contents), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=auth_headers,
    )

    assert confirm_res.status_code == 200, f"Confirm failed: {confirm_res.text}"
    body = confirm_res.json()
    assert body["message"] == "XLSX ingestion successful"
    assert body["validation_report"]["accepted_rows"] == 6
    assert body["validation_report"]["document_count"] == 3
    with TestingSessionLocal() as session:
        assert session.scalar(select(func.count()).select_from(JournalEntry)) == 3
        assert session.scalar(select(func.count()).select_from(JournalLine)) == 6

    run_res = client.post(
        f"/entities/{test_entity}/scrutiny-run?period_start=2025-04-01&period_end=2026-03-31",
        headers=auth_headers,
    )
    assert run_res.status_code == 409
    assert run_res.json()["detail"]["readiness"] == "PARTIAL"


def test_xlsx_fail_loud_validations_blank_canonical_value(auth_headers, test_entity):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Document Number", "G/L Account", "Posting Date", "Amount in local currency"])
    ws.append(["DOC-1", "", date(2025, 4, 1), 100])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    res = client.post(
        f"/entities/{test_entity}/upload-xlsx/confirm",
        data={
            "column_mapping": "{}",
            "target_period_start": "2025-04-01",
            "target_period_end": "2026-03-31"
        },
        files={"file": ("blank_ledger.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=auth_headers
    )

    assert res.status_code == 400, f"Expected 400, got: {res.status_code} {res.text}"
    detail = res.json()["detail"]
    assert detail["failed_batch_status"] == "FAILED"
    assert detail["validation_report"]["errors"]


def test_xlsx_fail_loud_validations_unbalanced_document(auth_headers, test_entity):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Document Number", "G/L Account", "Posting Date", "Amount in local currency"])
    ws.append(["DOC-1", "1000", date(2025, 4, 1), 100])
    ws.append(["DOC-1", "2000", date(2025, 4, 1), -99])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    res = client.post(
        f"/entities/{test_entity}/upload-xlsx/confirm",
        data={
            "column_mapping": "{}",
            "target_period_start": "2025-04-01",
            "target_period_end": "2026-03-31"
        },
        files={"file": ("unbalanced.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=auth_headers
    )

    assert res.status_code == 400, f"Expected 400, got: {res.status_code} {res.text}"
    detail = res.json()["detail"]
    assert detail["validation_report"]["errors"]


def test_xlsx_api_does_not_reflect_sensitive_document_number_in_errors(auth_headers, test_entity):
    document_number = "https://user:pass@example.test/path?query=opaque#fragment"
    contents = _xlsx_bytes(
        ["Document Number", "G/L Account", "Posting Date", "Amount in local currency"],
        [
            [document_number, "1000", date(2025, 4, 1), 100],
            [document_number, "2000", date(2025, 4, 1), -99],
        ],
    )

    response = client.post(
        f"/entities/{test_entity}/upload-xlsx/confirm",
        data={
            "column_mapping": "{}",
            "target_period_start": "2025-04-01",
            "target_period_end": "2026-03-31",
        },
        files={
            "file": (
                "sensitive-document.xlsx",
                io.BytesIO(contents),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=auth_headers,
    )

    assert response.status_code == 400, response.text
    detail = response.json()["detail"]
    for value in (document_number, "https://", "user:pass", "?query=opaque", "#fragment"):
        assert value not in str(detail)
        assert value not in str(detail["validation_report"])
        assert value not in str(detail["errors"])
    assert "balance" in detail["message"].lower()


def test_xlsx_fixed_profile_preserves_signed_rows(auth_headers, test_entity):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Document Number", "G/L Account", "Posting Date", "Amount in local currency"])
    ws.append(["DOC-1", "1000", date(2025, 4, 1), 100])
    ws.append(["DOC-1", "2000", date(2025, 4, 1), -100])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    res = client.post(
        f"/entities/{test_entity}/upload-xlsx/confirm",
        data={
            "column_mapping": "{}",
            "target_period_start": "2025-04-01",
            "target_period_end": "2026-03-31"
        },
        files={"file": ("pos_credit.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=auth_headers
    )
    assert res.status_code == 200, f"Fixed profile test failed: {res.status_code} {res.text}"
    assert Decimal(res.json()["validation_report"]["debit_total"]) == Decimal("100.00")
    assert Decimal(res.json()["validation_report"]["credit_total"]) == Decimal("100.00")

def test_xlsx_skips_blank_lines(auth_headers, test_entity):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Document Number", "G/L Account", "Posting Date", "Amount in local currency"])
    ws.append(["DOC-1", "1000", date(2025, 4, 1), 100])
    ws.append(["DOC-1", "2000", date(2025, 4, 1), -100])
    ws.append(["", "", "", ""])  # blank row
    ws.append(["DOC-2", "1000", date(2025, 4, 2), 200])
    ws.append(["DOC-2", "2000", date(2025, 4, 2), -200])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    res = client.post(
        f"/entities/{test_entity}/upload-xlsx/confirm",
        data={
            "column_mapping": "{}",
            "target_period_start": "2025-04-01",
            "target_period_end": "2026-03-31"
        },
        files={"file": ("blank_lines.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=auth_headers
    )
    assert res.status_code == 200, f"Failed to ingest: {res.text}"

    assert res.json()["validation_report"]["accepted_rows"] == 4
    assert res.json()["validation_report"]["skipped_rows"] == 1

    # A canonical journal without a baseline is not ready for scrutiny.
    run_res = client.post(
        f"/entities/{test_entity}/scrutiny-run?period_start=2025-04-01&period_end=2026-03-31",
        headers=auth_headers
    )
    assert run_res.status_code == 409, f"Scrutiny run unexpectedly succeeded: {run_res.text}"
    assert run_res.json()["detail"]["readiness"] == "PARTIAL"
