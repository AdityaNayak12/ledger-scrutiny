"""Fixed canonical GL XLSX ingestion and replay validation."""

import io
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from typing import Any, Dict, Mapping, Optional

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models import (
    Entity,
    ImportBatch,
    JournalEntry,
    JournalLine,
    LedgerAccount,
)
from app.ingestion.batches import fail_import_batch
from app.ingestion.reconciliation import build_reconciliation_report
from app.ingestion.schema import (
    BatchKind,
    DOCUMENT_BALANCE_TOLERANCE,
    GL_REQUIRED_HEADERS,
    JournalEntryRecord,
    JournalLineRecord,
    JournalLineSide,
    SourceFamily,
)


_OPTIONAL_ALIASES = {
    "document_type": ("Document Type",),
    "document_date": ("Document Date",),
    "posting_key": ("Posting Key",),
    "reference": (
        "Reference", "Invoice Reference", "Invoice/Reference", "Invoice / Reference", "Invoice",
        "Invoice No/Reference",
    ),
    "clearing_document": ("Clearing Document",),
    "profit_center": ("Profit Centre", "Profit Center"),
    "cost_center": ("Cost Centre", "Cost Center"),
    "text": ("Text", "Description", "Text/Bid/Cont/TndrNo", "L DESCRIPTION"),
    "supplier": ("Supplier", "Vendor", "Supplier/Vendor", "Vendor Name"),
    "wbs": ("WBS", "WBS Element", "WBS element"),
    "purchasing_document": ("Purchasing Document",),
    "customer": ("Customer", "Customer Name"),
    "quantity": ("Quantity",),
    "currency": ("Currency", "Currency Code"),
}
class _ParsedGL:
    __slots__ = ("entries", "input_rows", "skipped_rows", "skip_reasons", "warnings", "coverage_start", "coverage_end")

    def __init__(
        self,
        entries: tuple[JournalEntryRecord, ...],
        input_rows: int,
        skipped_rows: int,
        skip_reasons: tuple[dict[str, Any], ...],
        warnings: tuple[str, ...],
        coverage_start: date,
        coverage_end: date,
    ) -> None:
        self.entries = entries
        self.input_rows = input_rows
        self.skipped_rows = skipped_rows
        self.skip_reasons = skip_reasons
        self.warnings = warnings
        self.coverage_start = coverage_start
        self.coverage_end = coverage_end


class _GLParseError(ValueError):
    """Validation error carrying the reconciliation progress available before failure."""

    def __init__(
        self,
        message: str,
        *,
        input_rows: int = 0,
        skipped_rows: int = 0,
        skip_reasons: tuple[dict[str, Any], ...] = (),
        header_row: int = 0,
    ) -> None:
        super().__init__(message)
        self.input_rows = input_rows
        self.skipped_rows = skipped_rows
        self.skip_reasons = skip_reasons
        self.header_row = header_row


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text_value(value: Any) -> Optional[str]:
    if _is_blank(value):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _parse_date(value: Any, row_number: int, header: str, *, required: bool) -> Optional[date]:
    if _is_blank(value):
        if required:
            raise ValueError(f"Row {row_number}: {header} is a required value.")
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        candidate = value.strip()
        try:
            return date.fromisoformat(candidate)
        except ValueError:
            try:
                return datetime.fromisoformat(candidate).date()
            except ValueError:
                pass
    raise ValueError(f"Row {row_number}: {header} must be a valid date; got {value!r}.")


def _parse_decimal(value: Any, row_number: int, header: str, *, required: bool) -> Optional[Decimal]:
    if _is_blank(value):
        if required:
            raise ValueError(f"Row {row_number}: {header} is a required value.")
        return None
    if isinstance(value, bool):
        raise ValueError(f"Row {row_number}: {header} must be a valid decimal; got {value!r}.")
    try:
        result = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, TypeError, ValueError):
        raise ValueError(f"Row {row_number}: {header} must be a valid decimal; got {value!r}.") from None
    if not result.is_finite():
        raise ValueError(f"Row {row_number}: {header} must be a valid decimal; got {value!r}.")
    return result


def _validate_numeric_capacity(
    value: Decimal,
    row_number: int,
    header: str,
    *,
    scale: int,
    precision: int = 20,
) -> Decimal:
    decimal_tuple = value.as_tuple()
    decimal_scale = max(0, -decimal_tuple.exponent)
    if decimal_scale > scale:
        raise ValueError(
            f"Row {row_number}: {header} supports at most {scale} decimal places "
            f"(Numeric({precision},{scale})); got {value}."
        )
    integer_digits = max(0, value.adjusted() + 1) if value else 0
    if integer_digits > precision - scale:
        raise ValueError(
            f"Row {row_number}: {header} exceeds Numeric({precision},{scale}) capacity; got {value}."
        )
    return value


def _parse_storage_decimal(
    value: Any,
    row_number: int,
    header: str,
    *,
    required: bool,
    scale: int,
) -> Optional[Decimal]:
    result = _parse_decimal(value, row_number, header, required=required)
    if result is None:
        return None
    return _validate_numeric_capacity(result, row_number, header, scale=scale)


def _effective_coverage(
    target_period_start: Any,
    target_period_end: Any,
    *,
    declared_coverage: tuple[Any, Any] | None = None,
    declared_coverage_start: Any = None,
    declared_coverage_end: Any = None,
    coverage_start: Any = None,
    coverage_end: Any = None,
) -> tuple[date, date]:
    if declared_coverage is not None:
        declared_coverage_start, declared_coverage_end = declared_coverage
    start_value = declared_coverage_start if declared_coverage_start is not None else coverage_start
    end_value = declared_coverage_end if declared_coverage_end is not None else coverage_end
    start_value = target_period_start if start_value is None else start_value
    end_value = target_period_end if end_value is None else end_value
    start = _parse_date(start_value, 0, "Declared coverage start", required=True)
    end = _parse_date(end_value, 0, "Declared coverage end", required=True)
    assert start is not None and end is not None
    if start > end:
        raise ValueError(f"Declared coverage start {start} is after end {end}.")
    return start, end


def _header_name(value: Any, column_number: int) -> str:
    return str(value).strip() if not _is_blank(value) else f"__column_{column_number}"


def _fixed_header_row(worksheet: Any) -> tuple[int, list[str], dict[str, int]]:
    max_scan_row = min(15, worksheet.max_row)
    for row_number in range(1, max_scan_row + 1):
        values = [cell.value for cell in worksheet[row_number]]
        headers = [_header_name(value, index + 1) for index, value in enumerate(values)]
        positions: dict[str, int] = {}
        for index, value in enumerate(values):
            if value in GL_REQUIRED_HEADERS:
                if value in positions:
                    raise ValueError(f"Required header {value!r} appears more than once.")
                positions[value] = index
        if all(header in positions for header in GL_REQUIRED_HEADERS):
            return row_number, headers, positions
    raise ValueError(
        "Could not locate the exact required GL headers in the first 15 rows. "
        f"Required headers: {', '.join(GL_REQUIRED_HEADERS)}."
    )


def _optional_columns(headers: list[str]) -> dict[str, tuple[str, int]]:
    columns: dict[str, tuple[str, int]] = {}
    for field_name, aliases in _OPTIONAL_ALIASES.items():
        for alias in aliases:
            if alias in headers:
                columns[field_name] = (alias, headers.index(alias))
                break
    return columns


def _formula_cells(file_bytes: bytes, worksheet_title: str, header_row: int, required_columns: dict[str, int]) -> set[tuple[int, int]]:
    """Return original formula cells outside the required fields without evaluating them."""
    try:
        formula_workbook = openpyxl.load_workbook(
            io.BytesIO(file_bytes), data_only=False, read_only=True
        )
    except Exception as error:
        raise ValueError(f"Failed to inspect original Excel formula cells: {error}") from error
    try:
        formula_sheet = formula_workbook[worksheet_title]
        required_indexes = set(required_columns.values())
        formulas: set[tuple[int, int]] = set()
        for row_number, row in enumerate(formula_sheet.iter_rows(min_row=header_row + 1), header_row + 1):
            for column_index, cell in enumerate(row):
                if column_index in required_indexes:
                    continue
                value = cell.value
                if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                    formulas.add((row_number, column_index))
        return formulas
    finally:
        formula_workbook.close()


def _summary_skip_reason(
    values: list[Any],
    headers: list[str],
    required_columns: dict[str, int],
    row_number: int,
) -> dict[str, Any] | None:
    """Recognize explicit zero-valued footer labels, not arbitrary malformed rows."""
    required_indexes = set(required_columns.values())
    marker: tuple[str, str] | None = None
    for index, value in enumerate(values):
        if index in required_indexes or _is_blank(value):
            continue
        text = _text_value(value)
        normalized = text.upper() if text else ""
        if normalized == "TOTAL" or normalized.endswith("TOTAL") or normalized.endswith("SUBTOTAL"):
            marker = (headers[index], text)
            break
    if marker is None:
        return None

    identity_headers = ("Document Number", "G/L Account", "Posting Date")
    if any(not _is_blank(values[required_columns[header]]) for header in identity_headers):
        return None
    amount_value = values[required_columns["Amount in local currency"]]
    if _is_blank(amount_value):
        return None
    try:
        amount = Decimal(str(amount_value).replace(",", "").strip())
    except (InvalidOperation, TypeError, ValueError):
        return None
    if not amount.is_finite() or amount != 0:
        return None
    return {
        "row": row_number,
        "reason": "summary/footer row",
        "column": marker[0],
        "marker": marker[1],
    }


def _warning_messages(
    warning_counts: Counter[tuple[str, str]],
    invalid_quantity_counts: Counter[str],
) -> tuple[str, ...]:
    messages = [
        f"{category} for optional column '{header}' in {count} row(s)."
        for (category, header), count in sorted(warning_counts.items())
    ]
    messages.extend(
        f"optional column '{header}' contains non-numeric values in {count} row(s); "
        "canonical quantity left unset."
        for header, count in sorted(invalid_quantity_counts.items())
    )
    return tuple(messages)


def _failure_reconciliation_report(file_bytes: bytes, error: ValueError) -> dict[str, Any]:
    """Build a JSON-safe failure report even when parsing stops at the first bad row."""
    input_rows = getattr(error, "input_rows", 0)
    skipped_rows = getattr(error, "skipped_rows", 0)
    skip_reasons = list(getattr(error, "skip_reasons", ()))
    header_row = getattr(error, "header_row", 0)
    if not input_rows or not header_row:
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
            try:
                worksheet = workbook.active
                if worksheet is not None and worksheet.max_row:
                    header_row, _headers, _required_columns = _fixed_header_row(worksheet)
                    input_rows = max(0, worksheet.max_row - header_row)
            finally:
                workbook.close()
        except Exception:
            pass

    row_match = re.search(r"Row (\d+):", str(error))
    safe_error = safe_xlsx_error(error)
    failed_row = int(row_match.group(1)) if row_match else None
    skipped_row_numbers = {
        int(reason["row"])
        for reason in skip_reasons
        if isinstance(reason.get("row"), int)
    }
    reject_reasons: list[dict[str, Any]] = []
    if failed_row is not None:
        reject_reasons.append({"reason": safe_error, "row": failed_row})
    if header_row:
        for row_number in range(header_row + 1, header_row + input_rows + 1):
            if row_number in skipped_row_numbers or row_number == failed_row:
                continue
            reason = safe_error if failed_row is None else (
                "not processed because an earlier hard validation failure stopped parsing."
                if row_number > failed_row
                else "import rejected because the batch had a hard validation failure."
            )
            reject_reasons.append({"row": row_number, "reason": reason})
    elif failed_row is None:
        reject_reasons.append({"reason": safe_error})
    rejected_rows = max(0, input_rows - skipped_rows)
    report = build_reconciliation_report(
        (), input_rows=input_rows, rejected=rejected_rows, errors=[safe_error], coverage_complete=False,
    )
    report.update({
        "parser": "xlsx_gl",
        "source_family": SourceFamily.GL_UPLOAD.value,
        "accepted": 0,
        "accepted_rows": 0,
        "skipped": skipped_rows,
        "skipped_rows": skipped_rows,
        "skip_reasons": skip_reasons,
        "reject_reasons": reject_reasons,
    })
    return report


def _parse_fixed_gl_xlsx(
    file_bytes: bytes,
    target_period_start: Any,
    target_period_end: Any,
    *,
    declared_coverage: tuple[Any, Any] | None = None,
    declared_coverage_start: Any = None,
    declared_coverage_end: Any = None,
    coverage_start: Any = None,
    coverage_end: Any = None,
) -> _ParsedGL:
    period_start, period_end = _effective_coverage(
        target_period_start,
        target_period_end,
        declared_coverage=declared_coverage,
        declared_coverage_start=declared_coverage_start,
        declared_coverage_end=declared_coverage_end,
        coverage_start=coverage_start,
        coverage_end=coverage_end,
    )
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as error:
        raise ValueError(f"Failed to parse Excel file: {error}") from error

    input_rows = 0
    skip_reasons: list[dict[str, Any]] = []
    header_row = 0
    try:
        worksheet = workbook.active
        if worksheet is None or worksheet.max_row == 0:
            raise ValueError("Workbook has no active sheet or the active sheet is empty.")
        header_row, headers, required_columns = _fixed_header_row(worksheet)
        optional_columns = _optional_columns(headers)
        formula_cells = _formula_cells(file_bytes, worksheet.title, header_row, required_columns)
        entries: dict[str, dict[str, Any]] = {}
        warning_counts: Counter[tuple[str, str]] = Counter()
        invalid_quantity_counts: Counter[str] = Counter()
        invalid_quantity_headers: set[str] = set()
        input_rows = max(0, worksheet.max_row - header_row)

        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, max_row=worksheet.max_row, values_only=True),
            header_row + 1,
        ):
            values = list(row)
            if all(_is_blank(value) for value in values):
                skip_reasons.append({"row": row_number, "reason": "blank source row"})
                continue

            summary_reason = _summary_skip_reason(values, headers, required_columns, row_number)
            if summary_reason is not None:
                skip_reasons.append(summary_reason)
                continue

            def required_value(header: str) -> Any:
                value = values[required_columns[header]] if required_columns[header] < len(values) else None
                if _is_blank(value):
                    raise ValueError(f"Row {row_number}: {header} is a required value.")
                return value

            source_document_id = _text_value(required_value("Document Number"))
            account_code = _text_value(required_value("G/L Account"))
            posting_date = _parse_date(
                required_value("Posting Date"), row_number, "Posting Date", required=True
            )
            amount = _parse_decimal(
                required_value("Amount in local currency"),
                row_number,
                "Amount in local currency",
                required=True,
            )
            assert source_document_id is not None
            assert account_code is not None and posting_date is not None and amount is not None
            amount = _validate_numeric_capacity(
                amount, row_number, "Amount in local currency", scale=2
            )
            if not period_start <= posting_date <= period_end:
                raise ValueError(
                    f"Row {row_number}: posting date {posting_date} is outside declared coverage "
                    f"{period_start} to {period_end}."
                )

            side = JournalLineSide.DEBIT if amount >= 0 else JournalLineSide.CREDIT

            def optional_value(field_name: str) -> Any:
                if field_name not in optional_columns:
                    return None
                _header, column_index = optional_columns[field_name]
                return values[column_index] if column_index < len(values) else None

            def optional_text(field_name: str) -> Optional[str]:
                return _text_value(optional_value(field_name))

            document_date = _parse_date(
                optional_value("document_date"), row_number, "Document Date", required=False
            )
            quantity = None
            quantity_header = optional_columns.get("quantity", (None, 0))[0]
            raw_quantity = optional_value("quantity")
            if not _is_blank(raw_quantity):
                try:
                    quantity = _parse_storage_decimal(
                        raw_quantity, row_number, "Quantity", required=False, scale=4
                    )
                except ValueError as error:
                    if "must be a valid decimal" not in str(error):
                        raise
                    assert quantity_header is not None
                    invalid_quantity_headers.add(quantity_header)
                    invalid_quantity_counts[quantity_header] += 1
            document_type = optional_text("document_type")
            text = optional_text("text")

            for column_index, header in enumerate(headers):
                if header in GL_REQUIRED_HEADERS:
                    continue
                value = values[column_index] if column_index < len(values) else None
                if _is_blank(value):
                    category = (
                        "missing cached formula value"
                        if (row_number, column_index) in formula_cells
                        else "optional dimension is blank"
                    )
                    warning_counts[(category, header)] += 1

            source_metadata: dict[str, Any] = {}
            key_counts: defaultdict[str, int] = defaultdict(int)
            source_keys: list[str] = []
            for index, header in enumerate(headers):
                key_counts[header] += 1
                source_key = header if key_counts[header] == 1 else f"{header}#{key_counts[header]}"
                source_keys.append(source_key)
                source_metadata[source_key] = _json_value(values[index] if index < len(values) else None)

            known_optional_headers = {header for header, _index in optional_columns.values()}
            dimensions = {
                source_keys[index]: source_metadata[source_keys[index]]
                for index, header in enumerate(headers)
                if header not in GL_REQUIRED_HEADERS
                and (header not in known_optional_headers or header in invalid_quantity_headers)
            }
            line = JournalLineRecord(
                source_row_number=row_number,
                ledger_account_code=account_code,
                amount=amount,
                side=side,
                posting_key=optional_text("posting_key"),
                quantity=quantity,
                currency=optional_text("currency"),
                reference=optional_text("reference"),
                clearing_document=optional_text("clearing_document"),
                profit_center=optional_text("profit_center"),
                cost_center=optional_text("cost_center"),
                text=text,
                supplier=optional_text("supplier"),
                wbs=optional_text("wbs"),
                purchasing_document=optional_text("purchasing_document"),
                customer=optional_text("customer"),
                dimensions=dimensions,
                source_metadata=source_metadata,
            )
            document = entries.setdefault(
                source_document_id,
                {
                    "posting_date": posting_date,
                    "document_date": document_date,
                    "document_type": document_type,
                    "narration": text,
                    "lines": [],
                },
            )
            if document["posting_date"] != posting_date:
                raise ValueError(
                    f"Document {source_document_id} has inconsistent Posting Date values "
                    f"({document['posting_date']} and {posting_date})."
                )
            if document["document_date"] is None and document_date is not None:
                document["document_date"] = document_date
            if document["document_type"] is None and document_type is not None:
                document["document_type"] = document_type
            if document["narration"] is None and text is not None:
                document["narration"] = text
            document["lines"].append(line)

        canonical_entries = tuple(
            JournalEntryRecord(
                source_document_id=document_id,
                posting_date=document["posting_date"],
                document_date=document["document_date"],
                document_type=document["document_type"],
                narration=document["narration"],
                lines=tuple(document["lines"]),
                source_family=SourceFamily.GL_UPLOAD,
            )
            for document_id, document in entries.items()
        )
        for entry in canonical_entries:
            signed_total = sum((line.amount for line in entry.lines), Decimal("0"))
            if abs(signed_total) > DOCUMENT_BALANCE_TOLERANCE:
                raise ValueError(
                    f"Document {entry.source_document_id} does not balance within "
                    f"₹{DOCUMENT_BALANCE_TOLERANCE}: signed balance is {signed_total}."
                )
        return _ParsedGL(
            entries=canonical_entries,
            input_rows=input_rows,
            skipped_rows=len(skip_reasons),
            skip_reasons=tuple(skip_reasons),
            warnings=_warning_messages(warning_counts, invalid_quantity_counts),
            coverage_start=period_start,
            coverage_end=period_end,
        )
    except ValueError as error:
        if isinstance(error, _GLParseError):
            raise
        raise _GLParseError(
            str(error),
            input_rows=input_rows,
            skipped_rows=len(skip_reasons),
            skip_reasons=tuple(skip_reasons),
            header_row=header_row,
        ) from error
    finally:
        workbook.close()


def parse_gl_xlsx(
    file_bytes: bytes,
    target_period_start: Any,
    target_period_end: Any,
    *,
    declared_coverage: tuple[Any, Any] | None = None,
    declared_coverage_start: Any = None,
    declared_coverage_end: Any = None,
    coverage_start: Any = None,
    coverage_end: Any = None,
) -> tuple[JournalEntryRecord, ...]:
    """Parse the fixed GL profile into source-preserving canonical records."""
    return _parse_fixed_gl_xlsx(
        file_bytes,
        target_period_start,
        target_period_end,
        declared_coverage=declared_coverage,
        declared_coverage_start=declared_coverage_start,
        declared_coverage_end=declared_coverage_end,
        coverage_start=coverage_start,
        coverage_end=coverage_end,
    ).entries


def _entity(session: Session, entity_id: int) -> Entity:
    entity = session.execute(select(Entity).where(Entity.id == entity_id)).scalar_one_or_none()
    if entity is None:
        raise ValueError(f"Entity with ID {entity_id} not found.")
    return entity


def safe_xlsx_error(error: BaseException) -> str:
    """Keep workbook/parser failures actionable without exposing library details."""
    detail = str(error)
    row_match = re.fullmatch(r"(Row \d+): (.+)", detail)
    if row_match:
        row_label, row_detail = row_match.groups()
        if row_detail.endswith("is a required value."):
            return f"{row_label}: {row_detail}"
        valid_value = re.match(r"(.+?) must be a valid (date|decimal)", row_detail)
        if valid_value:
            return f"{row_label}: {valid_value.group(1)} must be a valid {valid_value.group(2)}."
        if "outside declared coverage" in row_detail:
            return f"{row_label}: posting date is outside declared coverage."
        if "numeric capacity" in row_detail:
            return f"{row_label}: amount exceeds canonical numeric capacity."
        return f"{row_label}: invalid canonical row data."
    if detail.startswith("Document "):
        if "does not balance" in detail:
            return "XLSX import failed validation: document is unbalanced."
        if "inconsistent Posting Date" in detail:
            return "XLSX import failed validation: document has inconsistent posting dates."
        return "XLSX import failed validation: invalid document data."
    if detail.startswith((
        "Could not locate the exact required GL headers",
        "Workbook has no active sheet",
        "Declared coverage ",
        "Only the fixed canonical GL profile is supported",
        "Exact duplicate XLSX import",
    )):
        return detail
    return "XLSX import failed validation. Verify the canonical GL headers, dates, amounts, and balanced documents."


def _canonical_replay_complete(
    session: Session,
    *,
    batch_id: int,
    entity_id: int,
    period_start: date,
    period_end: date,
    entries: tuple[JournalEntryRecord, ...],
    accounts_by_code: Mapping[str, LedgerAccount],
) -> bool:
    """Verify canonical document/line identity before treating a replay as idempotent."""
    expected_entries = {entry.source_document_id: entry for entry in entries}
    if len(expected_entries) != len(entries):
        return False
    actual_entries = session.scalars(
        select(JournalEntry).where(JournalEntry.import_batch_id == batch_id)
    ).all()
    actual_entry_map = {entry.source_document_id: entry for entry in actual_entries}
    if len(actual_entry_map) != len(actual_entries) or actual_entry_map.keys() != expected_entries.keys():
        return False
    for source_document_id, expected in expected_entries.items():
        actual = actual_entry_map[source_document_id]
        if (
            actual.entity_id != entity_id
            or not period_start <= actual.posting_date <= period_end
            or actual.posting_date != expected.posting_date
            or actual.document_date != expected.document_date
            or actual.document_type != expected.document_type
            or actual.narration != expected.narration
        ):
            return False

    line_rows = session.execute(
        select(JournalLine, JournalEntry)
        .join(JournalEntry, JournalEntry.id == JournalLine.journal_entry_id)
        .where(JournalEntry.import_batch_id == batch_id)
    ).all()
    actual_line_map = {
        (entry.source_document_id, line.source_row_number): (line, entry)
        for line, entry in line_rows
    }
    if len(actual_line_map) != len(line_rows):
        return False

    expected_line_map: dict[tuple[str, int], JournalLineRecord] = {}
    expected_accounts: dict[str, LedgerAccount] = {}
    for entry in entries:
        for line in entry.lines:
            account = accounts_by_code.get(line.ledger_account_code)
            if account is None or account.entity_id != entity_id:
                return False
            key = (entry.source_document_id, line.source_row_number)
            if key in expected_line_map:
                return False
            expected_line_map[key] = line
            expected_accounts[line.ledger_account_code] = account
    if actual_line_map.keys() != expected_line_map.keys():
        return False

    actual_account_ids = {line.ledger_account_id for line, _entry in line_rows}
    actual_accounts = session.scalars(
        select(LedgerAccount).where(LedgerAccount.id.in_(actual_account_ids or {-1}))
    ).all()
    accounts_by_id = {account.id: account for account in actual_accounts}
    if len(accounts_by_id) != len(actual_account_ids) or any(
        account.entity_id != entity_id for account in actual_accounts
    ):
        return False

    def line_facts(line: JournalLine) -> tuple[Any, ...]:
        return (
            line.source_row_number,
            line.ledger_account_id,
            Decimal(str(line.amount)),
            line.side,
            line.posting_key,
            Decimal(str(line.quantity)) if line.quantity is not None else None,
            line.currency,
            line.reference,
            line.clearing_document,
            line.profit_center,
            line.cost_center,
            line.text,
            line.supplier,
            line.wbs,
            line.purchasing_document,
            line.customer,
            dict(line.dimensions or {}),
            dict(line.source_metadata or {}),
        )

    def expected_line_facts(line: JournalLineRecord) -> tuple[Any, ...]:
        account = expected_accounts[line.ledger_account_code]
        return (
            line.source_row_number,
            account.id,
            line.amount,
            line.side.value,
            line.posting_key,
            line.quantity,
            line.currency,
            line.reference,
            line.clearing_document,
            line.profit_center,
            line.cost_center,
            line.text,
            line.supplier,
            line.wbs,
            line.purchasing_document,
            line.customer,
            dict(line.dimensions),
            dict(line.source_metadata),
        )

    for key, expected in expected_line_map.items():
        actual, entry = actual_line_map[key]
        if entry.entity_id != entity_id or not period_start <= entry.posting_date <= period_end:
            return False
        if line_facts(actual) != expected_line_facts(expected):
            return False
    return True


def validate_gl_xlsx_replay(
    file_bytes: bytes,
    target_period_start: Any,
    target_period_end: Any,
    entity_id: int,
    session: Session,
    *,
    import_batch_id: int,
) -> dict[str, Any]:
    """Validate an exact active GL replay without writing another canonical row."""
    batch = session.get(ImportBatch, import_batch_id)
    if batch is None or batch.entity_id != entity_id or batch.status != "ACTIVE":
        raise ValueError("Exact duplicate XLSX import has no active canonical batch.")
    if batch.source_family not in (None, SourceFamily.GL_UPLOAD.value):
        raise ValueError("Exact duplicate XLSX import is not a canonical GL batch.")
    period_start, period_end = _effective_coverage(target_period_start, target_period_end)
    if (batch.coverage_start, batch.coverage_end) != (period_start, period_end):
        raise ValueError("Exact duplicate XLSX import does not match the active batch period.")
    parsed = _parse_fixed_gl_xlsx(file_bytes, period_start, period_end)
    accounts = {
        account.external_code: account
        for account in session.scalars(select(LedgerAccount).where(LedgerAccount.entity_id == entity_id)).all()
        if account.external_code
    }
    if not _canonical_replay_complete(
        session,
        batch_id=import_batch_id,
        entity_id=entity_id,
        period_start=period_start,
        period_end=period_end,
        entries=parsed.entries,
        accounts_by_code=accounts,
    ):
        raise ValueError("Exact duplicate XLSX import contains missing or altered canonical journal children.")
    return dict(batch.validation_report or {})


def _fail_staged_gl_contract(session: Session, batch: ImportBatch, message: str) -> None:
    report = build_reconciliation_report(
        (), input_rows=0, errors=[message], coverage_complete=False,
    )
    report.update({
        "parser": "xlsx_gl",
        "source_family": SourceFamily.GL_UPLOAD.value,
        "reject_reasons": [],
    })
    fail_import_batch(session, batch, errors=[message], validation_report=report)
    raise ValueError(message)


def normalize_gl_xlsx(
    file_bytes: bytes,
    target_period_start: Any,
    target_period_end: Any,
    entity_id: int,
    session: Session,
    *,
    import_batch_id: Optional[int],
    declared_coverage: tuple[Any, Any] | None = None,
    declared_coverage_start: Any = None,
    declared_coverage_end: Any = None,
    coverage_start: Any = None,
    coverage_end: Any = None,
) -> dict[str, Any]:
    """Parse and persist fixed-profile GL records without depending on FastAPI."""
    entity = _entity(session, entity_id)
    if import_batch_id is None:
        raise ValueError("Fixed GL normalization requires an import_batch_id.")
    batch = session.get(ImportBatch, import_batch_id)
    if batch is None:
        raise ValueError(f"Import batch {import_batch_id} not found.")
    if batch.entity_id != entity_id:
        raise ValueError(f"Import batch {import_batch_id} does not belong to entity {entity_id}.")
    if batch.status != "STAGED":
        raise ValueError(
            f"Fixed GL normalization requires a STAGED ImportBatch; batch {import_batch_id} is {batch.status}."
        )
    if batch.kind != BatchKind.JOURNAL.value:
        _fail_staged_gl_contract(
            session,
            batch,
            "Fixed GL normalization requires a journal ImportBatch.",
        )
    if batch.source_family != SourceFamily.GL_UPLOAD.value:
        _fail_staged_gl_contract(
            session,
            batch,
            "Fixed GL normalization requires a gl_upload source_family.",
        )
    received_sha256 = sha256(bytes(file_bytes)).hexdigest()
    if received_sha256 != batch.content_sha256:
        _fail_staged_gl_contract(
            session,
            batch,
            "GL workbook bytes do not match staged ImportBatch content_sha256 "
            f"(staged {batch.content_sha256}, received {received_sha256}).",
        )

    if (
        declared_coverage is None
        and declared_coverage_start is None
        and declared_coverage_end is None
        and coverage_start is None
        and coverage_end is None
        and batch.coverage_start is not None
        and batch.coverage_end is not None
    ):
        declared_coverage = (batch.coverage_start, batch.coverage_end)

    try:
        parsed = _parse_fixed_gl_xlsx(
            file_bytes,
            target_period_start,
            target_period_end,
            declared_coverage=declared_coverage,
            declared_coverage_start=declared_coverage_start,
            declared_coverage_end=declared_coverage_end,
            coverage_start=coverage_start,
            coverage_end=coverage_end,
        )
    except ValueError as error:
        fail_import_batch(
            session,
            batch,
            errors=[safe_xlsx_error(error)],
            validation_report=_failure_reconciliation_report(file_bytes, error),
        )
        raise

    accounts = session.execute(select(LedgerAccount).where(LedgerAccount.entity_id == entity_id)).scalars().all()
    accounts_by_code = {account.external_code: account for account in accounts if account.external_code}
    account_names = {account.name for account in accounts}
    unclassified_codes: set[str] = set()
    account_by_code: dict[str, LedgerAccount] = {}
    for entry in parsed.entries:
        for line in entry.lines:
            code = line.ledger_account_code
            account = account_by_code.get(code) or accounts_by_code.get(code)
            if account is None:
                account_name = code
                if account_name in account_names:
                    account_name = f"G/L {code}"
                    suffix = 2
                    while account_name in account_names:
                        account_name = f"G/L {code} ({suffix})"
                        suffix += 1
                account = LedgerAccount(entity_id=entity.id, external_code=code, name=account_name)
                session.add(account)
                session.flush()
                account_names.add(account_name)
            account_by_code[code] = account
            if account.group_name is None or account.normal_balance is None:
                unclassified_codes.add(code)

    warnings = list(parsed.warnings)
    warnings.extend(
        f"Account {code}: classification could not be resolved; group and normal balance remain unset."
        for code in sorted(unclassified_codes)
    )
    report = build_reconciliation_report(
        parsed.entries,
        input_rows=parsed.input_rows,
        skipped=parsed.skipped_rows,
        warnings=warnings,
        unmapped_account_codes=unclassified_codes,
        coverage_complete=True,
    )
    report.update({
        "parser": "xlsx_gl",
        "source_family": SourceFamily.GL_UPLOAD.value,
        "coverage_start": parsed.coverage_start.isoformat(),
        "coverage_end": parsed.coverage_end.isoformat(),
        "skip_reasons": list(parsed.skip_reasons),
    })

    for entry_record in parsed.entries:
        entry = JournalEntry(
            import_batch_id=batch.id,
            entity_id=entity.id,
            source_document_id=entry_record.source_document_id,
            posting_date=entry_record.posting_date,
            document_date=entry_record.document_date,
            document_type=entry_record.document_type,
            narration=entry_record.narration,
        )
        session.add(entry)
        for line_record in entry_record.lines:
            account = account_by_code[line_record.ledger_account_code]
            session.add(JournalLine(
                journal_entry=entry,
                ledger_account=account,
                source_row_number=line_record.source_row_number,
                amount=line_record.amount,
                side=line_record.side.value,
                posting_key=line_record.posting_key,
                quantity=line_record.quantity,
                currency=line_record.currency,
                reference=line_record.reference,
                clearing_document=line_record.clearing_document,
                profit_center=line_record.profit_center,
                cost_center=line_record.cost_center,
                text=line_record.text,
                supplier=line_record.supplier,
                wbs=line_record.wbs,
                purchasing_document=line_record.purchasing_document,
                customer=line_record.customer,
                dimensions=dict(line_record.dimensions),
                source_metadata=dict(line_record.source_metadata),
            ))

    batch.source_metadata = {
        **(batch.source_metadata or {}),
        "parser": "xlsx_gl",
        "required_headers": list(GL_REQUIRED_HEADERS),
    }
    batch.validation_report = report
    session.flush()
    return report


def _uses_fixed_gl_profile(column_mapping: Optional[Dict[str, str]]) -> bool:
    if not column_mapping:
        return True
    values = {str(value).strip() for value in column_mapping.values()}
    return set(GL_REQUIRED_HEADERS).issubset(values)


def normalize_xlsx_confirm(
    file_bytes: bytes,
    column_mapping: Optional[Dict[str, str]],
    sign_convention: str,
    target_period_start: str,
    target_period_end: str,
    entity_id: int,
    session: Session,
    clear_only_period: bool = True,
    import_batch_id: Optional[int] = None,
) -> Entity:
    """Dispatch only the fixed canonical GL profile."""
    entity = _entity(session, entity_id)
    if not _uses_fixed_gl_profile(column_mapping):
        raise ValueError(
            "Only the fixed canonical GL profile is supported for XLSX ingestion. "
            f"Use exact headers: {', '.join(GL_REQUIRED_HEADERS)}."
        )
    if sign_convention != "negative_is_credit":
        raise ValueError(
            "Fixed canonical GL imports require signed amounts: positive values are debits and negative values are credits."
        )
    normalize_gl_xlsx(
        file_bytes,
        target_period_start,
        target_period_end,
        entity_id,
        session,
        import_batch_id=import_batch_id,
    )
    return entity
