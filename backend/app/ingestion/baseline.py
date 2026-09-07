"""Signed account-level closing-balance ingestion."""

import base64
import io
import json
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from typing import Any, Iterable, Mapping

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models import BalanceCheckpoint, Entity, ImportBatch, LedgerAccount
from app.ingestion.batches import activate_import_batch, fail_import_batch
from app.ingestion.reconciliation import build_reconciliation_report, compute_dataset_fingerprint
from app.ingestion.schema import (
    DOCUMENT_BALANCE_TOLERANCE,
    BalanceCheckpointRecord,
    BatchKind,
    SourceFamily,
)


BASELINE_REQUIRED_HEADERS = ("Account Code", "Balance Date", "Signed Balance", "Currency")
_DECIMAL_ZERO = Decimal("0")
_PDF_PREFIX = b"%PDF"


class BaselineValidationError(ValueError):
    """A structured baseline failed parsing or accounting validation."""

    def __init__(
        self,
        message: str,
        *,
        records: Iterable[BalanceCheckpointRecord] = (),
        input_rows: int = 0,
        skipped_rows: int = 0,
        skip_reasons: Iterable[Mapping[str, Any]] = (),
        missing_account_codes: Iterable[str] = (),
        unknown_account_codes: Iterable[str] = (),
    ) -> None:
        super().__init__(message)
        self.records = tuple(records)
        self.input_rows = input_rows
        self.skipped_rows = skipped_rows
        self.skip_reasons = tuple(dict(reason) for reason in skip_reasons)
        self.missing_account_codes = tuple(sorted(set(missing_account_codes)))
        self.unknown_account_codes = tuple(sorted(set(unknown_account_codes)))


@dataclass(frozen=True, slots=True)
class _ParsedBaseline:
    records: tuple[BalanceCheckpointRecord, ...]
    input_rows: int
    skipped_rows: int
    skip_reasons: tuple[dict[str, Any], ...]


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text_value(value: Any) -> str | None:
    if _is_blank(value):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_date(value: Any, row_number: int) -> date:
    if _is_blank(value):
        raise BaselineValidationError(f"Row {row_number}: Balance Date is a required value.")
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        candidate = value.strip()
        try:
            return date.fromisoformat(candidate)
        except ValueError:
            try:
                return datetime.fromisoformat(candidate).date()
            except ValueError:
                pass
    raise BaselineValidationError(
        f"Row {row_number}: Balance Date must be a valid date; got {value!r}."
    )


def _parse_decimal(value: Any, row_number: int) -> Decimal:
    if _is_blank(value):
        raise BaselineValidationError(f"Row {row_number}: Signed Balance is a required value.")
    if isinstance(value, bool):
        raise BaselineValidationError(
            f"Row {row_number}: Signed Balance must be a valid decimal; got {value!r}."
        )
    try:
        result = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, TypeError, ValueError):
        raise BaselineValidationError(
            f"Row {row_number}: Signed Balance must be a valid decimal; got {value!r}."
        ) from None
    if not result.is_finite():
        raise BaselineValidationError(
            f"Row {row_number}: Signed Balance must be a valid decimal; got {value!r}."
        )
    decimal_scale = max(0, -result.as_tuple().exponent)
    if decimal_scale > 2:
        raise BaselineValidationError(
            f"Row {row_number}: Signed Balance supports at most 2 decimal places; got {result}."
        )
    integer_digits = max(0, result.adjusted() + 1) if result else 0
    if integer_digits > 18:
        raise BaselineValidationError(
            f"Row {row_number}: Signed Balance exceeds Numeric(20,2) capacity; got {result}."
        )
    return result


def _header_row(worksheet: Any) -> tuple[int, dict[str, int]]:
    values = list(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    unexpected = [
        value
        for value in values
        if not _is_blank(value) and value not in BASELINE_REQUIRED_HEADERS
    ]
    if unexpected:
        raise BaselineValidationError(
            "Baseline headers must be the exact required headers in worksheet row 1; "
            "unexpected nonblank header(s): "
            + ", ".join(repr(value) for value in unexpected)
            + ". Required headers: "
            + ", ".join(BASELINE_REQUIRED_HEADERS)
            + "."
        )
    positions: dict[str, int] = {}
    for header in BASELINE_REQUIRED_HEADERS:
        matches = [index for index, value in enumerate(values) if value == header]
        if len(matches) > 1:
            raise BaselineValidationError(
                f"Required baseline header {header!r} appears more than once."
            )
        if matches:
            positions[header] = matches[0]
    if len(positions) == len(BASELINE_REQUIRED_HEADERS):
        return 1, positions
    raise BaselineValidationError(
        "Baseline headers must be the exact required headers in worksheet row 1. "
        f"Required headers: {', '.join(BASELINE_REQUIRED_HEADERS)}."
    )


def _normalise_codes(codes: Iterable[Any] | None) -> set[str] | None:
    if codes is None:
        return None
    if isinstance(codes, str):
        codes = (codes,)
    result = {_text_value(code) for code in codes}
    result.discard(None)
    return {str(code) for code in result}


def _normalise_currency(currency: Any) -> str | None:
    value = _text_value(currency)
    return value.upper() if value is not None else None


def _record_summary(records: Iterable[BalanceCheckpointRecord]) -> dict[str, Any]:
    records = tuple(records)
    codes = {record.ledger_account_code for record in records}
    dates = {record.balance_date for record in records}
    currencies = {record.currency for record in records if record.currency is not None}
    signed_total = sum((record.balance for record in records), _DECIMAL_ZERO)
    debit = sum((record.balance for record in records if record.balance > 0), _DECIMAL_ZERO)
    credit = sum((-record.balance for record in records if record.balance < 0), _DECIMAL_ZERO)
    return {
        "account_codes": sorted(codes),
        "account_count": len(codes),
        "balance_date": next(iter(dates)).isoformat() if len(dates) == 1 else None,
        "currency": next(iter(currencies)) if len(currencies) == 1 else None,
        "signed_total": format(signed_total, "f"),
        "debit": format(debit, "f"),
        "credit": format(credit, "f"),
    }


def _baseline_fingerprint(
    records: Iterable[BalanceCheckpointRecord], summary: Mapping[str, Any]
) -> str:
    account_balance_pairs = sorted(
        (
            record.ledger_account_code,
            format(record.balance, "f"),
        )
        for record in records
    )
    return compute_dataset_fingerprint(({
        "kind": BatchKind.BALANCE_CHECKPOINT.value,
        "balance_date": summary["balance_date"],
        "currency": summary["currency"],
        "account_balance_pairs": account_balance_pairs,
    },))


def _validate_records(
    records: Iterable[BalanceCheckpointRecord],
    *,
    expected_account_codes: Iterable[Any] | None = None,
    allowed_account_codes: Iterable[Any] | None = None,
    expected_currency: str | None = None,
    expected_balance_date: date | None = None,
) -> dict[str, Any]:
    records = tuple(records)
    summary = _record_summary(records)
    errors: list[str] = []
    codes = set(summary["account_codes"])
    dates = {record.balance_date for record in records}
    currencies = {record.currency for record in records}
    duplicates = sorted(
        code for code in codes
        if sum(record.ledger_account_code == code for record in records) > 1
    )
    if not records:
        errors.append("Baseline schedule contains no account rows.")
    if duplicates:
        errors.append(f"Duplicate account code(s): {', '.join(duplicates)}.")
    if len(dates) != 1:
        errors.append(
            f"Baseline must contain one balance date; found {len(dates)} distinct dates."
        )
    if any(record.currency is None for record in records):
        errors.append("Currency is a required value for every baseline row.")
    if len(currencies) != 1:
        errors.append(
            f"Baseline rows must use the same currency; found {len(currencies)} distinct currencies."
        )

    signed_total = sum((record.balance for record in records), _DECIMAL_ZERO)
    if abs(signed_total) > DOCUMENT_BALANCE_TOLERANCE:
        errors.append(
            "Baseline signed total must balance within "
            f"{DOCUMENT_BALANCE_TOLERANCE}; got {signed_total}."
        )

    expected_codes = _normalise_codes(expected_account_codes)
    allowed_codes = _normalise_codes(allowed_account_codes)
    if allowed_codes is None:
        allowed_codes = expected_codes
    unknown = sorted(codes - allowed_codes) if allowed_codes is not None else []
    missing = sorted(expected_codes - codes) if expected_codes is not None else []
    if unknown:
        errors.append(f"Unknown account code(s): {', '.join(unknown)}.")
    if missing:
        errors.append(f"Missing expected account code(s): {', '.join(missing)}.")

    actual_currency = next(iter(currencies), None)
    if expected_currency is not None:
        expected_currency = _normalise_currency(expected_currency)
        if actual_currency != expected_currency:
            errors.append(
                f"Baseline currency {actual_currency!r} does not match expected currency "
                f"{expected_currency!r}."
            )
    actual_date = next(iter(dates), None)
    if isinstance(expected_balance_date, datetime):
        expected_balance_date = expected_balance_date.date()
    if expected_balance_date is not None and actual_date != expected_balance_date:
        errors.append(
            f"Baseline balance date {actual_date!r} does not match expected date "
            f"{expected_balance_date!r}."
        )

    if errors:
        raise BaselineValidationError(
            " ".join(errors),
            records=records,
            missing_account_codes=missing,
            unknown_account_codes=unknown,
        )

    return {
        **summary,
        "complete": not missing,
        "expected_account_count": len(expected_codes) if expected_codes is not None else None,
        "missing_account_codes": missing,
        "unknown_account_codes": unknown,
    }


def _parse_baseline_xlsx(
    file_bytes: bytes,
    *,
    expected_account_codes: Iterable[Any] | None = None,
    allowed_account_codes: Iterable[Any] | None = None,
    expected_currency: str | None = None,
    expected_balance_date: date | None = None,
) -> _ParsedBaseline:
    file_bytes = bytes(file_bytes)
    if file_bytes.lstrip().startswith(_PDF_PREFIX):
        raise BaselineValidationError(
            "PDF-only baseline rejected; a structured XLSX account schedule is required."
        )
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as error:
        raise BaselineValidationError(f"Failed to parse baseline Excel file: {error}") from error

    input_rows = 0
    skipped_rows = 0
    skip_reasons: list[dict[str, Any]] = []
    records: list[BalanceCheckpointRecord] = []
    try:
        worksheet = workbook.active
        if worksheet is None or worksheet.max_row == 0:
            raise BaselineValidationError("Workbook has no active sheet or the active sheet is empty.")
        header_number, columns = _header_row(worksheet)
        input_rows = max(0, worksheet.max_row - header_number)
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=header_number + 1, max_row=worksheet.max_row, values_only=True),
            header_number + 1,
        ):
            values = list(row)
            if all(_is_blank(value) for value in values):
                skipped_rows += 1
                skip_reasons.append({"row": row_number, "reason": "blank source row"})
                continue
            account_code = _text_value(
                values[columns["Account Code"]] if columns["Account Code"] < len(values) else None
            )
            if account_code is not None and account_code.upper() in {"TOTAL", "TOTALS", "GRAND TOTAL"}:
                if all(
                    _is_blank(values[columns[field]] if columns[field] < len(values) else None)
                    for field in ("Balance Date", "Currency")
                ):
                    skipped_rows += 1
                    skip_reasons.append({"row": row_number, "reason": "summary/footer row"})
                    continue
            if account_code is None:
                raise BaselineValidationError(
                    f"Row {row_number}: Account Code is a required value."
                )
            balance_date = _parse_date(
                values[columns["Balance Date"]] if columns["Balance Date"] < len(values) else None,
                row_number,
            )
            balance = _parse_decimal(
                values[columns["Signed Balance"]] if columns["Signed Balance"] < len(values) else None,
                row_number,
            )
            currency = _normalise_currency(
                values[columns["Currency"]] if columns["Currency"] < len(values) else None
            )
            if currency is None:
                raise BaselineValidationError(
                    f"Row {row_number}: Currency is a required value."
                )
            records.append(BalanceCheckpointRecord(
                ledger_account_code=account_code,
                balance_date=balance_date,
                balance=balance,
                currency=currency,
                source_family=SourceFamily.GL_UPLOAD,
                source_metadata={"source_row_number": row_number},
            ))
        try:
            _validate_records(
                records,
                expected_account_codes=expected_account_codes,
                allowed_account_codes=allowed_account_codes,
                expected_currency=expected_currency,
                expected_balance_date=expected_balance_date,
            )
        except BaselineValidationError as error:
            error.input_rows = input_rows
            error.skipped_rows = skipped_rows
            error.skip_reasons = tuple(skip_reasons)
            raise
        return _ParsedBaseline(
            records=tuple(records),
            input_rows=input_rows,
            skipped_rows=skipped_rows,
            skip_reasons=tuple(skip_reasons),
        )
    except BaselineValidationError as error:
        if not error.input_rows:
            error.input_rows = input_rows
        if not error.skipped_rows:
            error.skipped_rows = skipped_rows
        if not error.skip_reasons:
            error.skip_reasons = tuple(skip_reasons)
        if not error.records:
            error.records = tuple(records)
        raise
    finally:
        workbook.close()


def parse_baseline_xlsx(
    file_bytes: bytes,
    *,
    expected_account_codes: Iterable[Any] | None = None,
    expected_currency: str | None = None,
    expected_balance_date: date | None = None,
) -> tuple[BalanceCheckpointRecord, ...]:
    """Parse and validate a signed account-level closing schedule."""
    return _parse_baseline_xlsx(
        file_bytes,
        expected_account_codes=expected_account_codes,
        expected_currency=expected_currency,
        expected_balance_date=expected_balance_date,
    ).records


def validate_baseline_records(
    records: Iterable[BalanceCheckpointRecord],
    *,
    expected_account_codes: Iterable[Any] | None = None,
    expected_currency: str | None = None,
    expected_balance_date: date | None = None,
) -> dict[str, Any]:
    """Validate already parsed records and return JSON-safe coverage data."""
    return _validate_records(
        records,
        expected_account_codes=expected_account_codes,
        expected_currency=expected_currency,
        expected_balance_date=expected_balance_date,
    )


def _baseline_report(
    records: Iterable[BalanceCheckpointRecord],
    *,
    input_rows: int,
    skipped_rows: int = 0,
    skip_reasons: Iterable[Mapping[str, Any]] = (),
    errors: Iterable[str] = (),
    expected_account_codes: Iterable[Any] | None = None,
    allowed_account_codes: Iterable[Any] | None = None,
) -> dict[str, Any]:
    records = tuple(records)
    errors = list(errors)
    summary = _record_summary(records)
    expected = _normalise_codes(expected_account_codes)
    allowed = _normalise_codes(allowed_account_codes)
    codes = set(summary["account_codes"])
    missing = sorted(expected - codes) if expected is not None else []
    unknown = sorted(codes - allowed) if allowed is not None else []
    coverage = {
        "present": bool(records),
        "complete": not errors and not missing and not unknown and bool(records),
        "account_count": summary["account_count"],
        "expected_account_count": len(expected) if expected is not None else None,
        "missing_account_codes": missing,
        "unknown_account_codes": unknown,
        "balance_date": summary["balance_date"],
        "currency": summary["currency"],
        "signed_total": summary["signed_total"],
    }
    report = build_reconciliation_report(
        (),
        baseline_coverage=coverage,
        input_rows=input_rows,
        skipped=skipped_rows,
        rejected=1 if errors else 0,
        errors=errors,
        coverage_complete=not errors and not missing and not unknown and bool(records),
        fingerprint=_baseline_fingerprint(records, summary),
    )
    report.update({
        "parser": "xlsx_baseline",
        "source_family": SourceFamily.GL_UPLOAD.value,
        "account_codes": summary["account_codes"],
        "account_count": summary["account_count"],
        "unmapped_account_codes": unknown,
        "unmapped_account_count": len(unknown),
        "balance_date": summary["balance_date"],
        "currency": summary["currency"],
        "signed_total": summary["signed_total"],
        "balance_total": summary["signed_total"],
        "debit": summary["debit"],
        "credit": summary["credit"],
        "net": summary["signed_total"],
        "debit_total": summary["debit"],
        "credit_total": summary["credit"],
        "net_total": summary["signed_total"],
        "accepted": len(records) if not errors else 0,
        "accepted_rows": len(records) if not errors else 0,
        "skipped": skipped_rows,
        "skipped_rows": skipped_rows,
        "skip_reasons": [dict(reason) for reason in skip_reasons],
        "expected_account_count": len(expected) if expected is not None else None,
        "missing_account_codes": missing,
        "unknown_account_codes": unknown,
        "reject_reasons": [{"reason": error} for error in errors],
    })
    return json.loads(json.dumps(report))


def _pdf_artifact(pdf_bytes: bytes, filename: str | None) -> dict[str, str]:
    pdf_bytes = bytes(pdf_bytes)
    return {
        "filename": filename or "signed-closing-baseline.pdf",
        "content_type": "application/pdf",
        "content_sha256": sha256(pdf_bytes).hexdigest(),
        "bytes_base64": base64.b64encode(pdf_bytes).decode("ascii"),
    }


def _retain_pdf_evidence(batch: ImportBatch, pdf_bytes: bytes, filename: str | None) -> None:
    metadata = dict(batch.source_metadata or {})
    evidence = dict(metadata.get("supporting_evidence") or {})
    evidence["signed_pdf"] = _pdf_artifact(pdf_bytes, filename)
    metadata["supporting_evidence"] = evidence
    batch.source_metadata = metadata


def _entity(session: Session, entity_id: int) -> Entity:
    entity = session.execute(select(Entity).where(Entity.id == entity_id)).scalar_one_or_none()
    if entity is None:
        raise ValueError(f"Entity with ID {entity_id} not found.")
    return entity


def normalize_baseline_xlsx(
    file_bytes: bytes,
    entity_id: int,
    session: Session,
    *,
    import_batch_id: int,
    expected_account_codes: Iterable[Any] | None = None,
    expected_currency: str | None = None,
    expected_balance_date: date | None = None,
    signed_pdf_bytes: bytes | None = None,
    signed_pdf_filename: str | None = None,
    activate: bool = True,
) -> dict[str, Any]:
    """Persist a validated baseline into a staged checkpoint batch."""
    entity = _entity(session, entity_id)
    batch = session.get(ImportBatch, import_batch_id)
    if batch is None:
        raise ValueError(f"Import batch {import_batch_id} not found.")
    if batch.entity_id != entity_id:
        raise ValueError(f"Import batch {import_batch_id} does not belong to entity {entity_id}.")
    if batch.kind != BatchKind.BALANCE_CHECKPOINT.value:
        raise ValueError(
            f"Baseline normalization requires a {BatchKind.BALANCE_CHECKPOINT.value} ImportBatch."
        )
    if batch.source_family != SourceFamily.GL_UPLOAD.value:
        raise ValueError(
            "Baseline normalization requires a gl_upload source_family for balance checkpoints."
        )
    accounts = session.execute(
        select(LedgerAccount).where(LedgerAccount.entity_id == entity_id)
    ).scalars().all()
    accounts_by_code = {
        str(account.external_code).strip(): account
        for account in accounts
        if account.external_code is not None and str(account.external_code).strip()
    }
    file_bytes = bytes(file_bytes)
    if batch.status == "ACTIVE" and sha256(file_bytes).hexdigest() == batch.content_sha256:
        parsed = _parse_baseline_xlsx(
            file_bytes,
            expected_account_codes=set(accounts_by_code),
            allowed_account_codes=set(accounts_by_code),
            expected_currency=expected_currency,
            expected_balance_date=expected_balance_date,
        )
        expected_checkpoints = {
            (accounts_by_code[record.ledger_account_code].id, record.balance_date): (
                record.balance,
                record.currency,
                {
                    **dict(record.source_metadata),
                    "source_filename": batch.original_filename,
                    "source_sha256": batch.content_sha256,
                },
            )
            for record in parsed.records
        }
        actual_checkpoints = session.execute(
            select(BalanceCheckpoint).where(BalanceCheckpoint.import_batch_id == batch.id)
        ).scalars().all()
        actual_by_key = {
            (checkpoint.ledger_account_id, checkpoint.balance_date): checkpoint
            for checkpoint in actual_checkpoints
        }
        if (
            len(actual_by_key) != len(actual_checkpoints)
            or actual_by_key.keys() != expected_checkpoints.keys()
            or any(
                (
                    checkpoint.balance,
                    checkpoint.currency,
                    dict(checkpoint.source_metadata or {}),
                ) != expected_checkpoints[key]
                for key, checkpoint in actual_by_key.items()
            )
        ):
            raise ValueError(
                "Exact duplicate baseline import contains missing or altered canonical baseline children."
            )
        return json.loads(json.dumps(batch.validation_report or {}))
    if batch.status != "STAGED":
        raise ValueError(
            f"Baseline normalization requires a STAGED ImportBatch; batch {import_batch_id} is {batch.status}."
        )
    evidence_bytes = signed_pdf_bytes
    if evidence_bytes is None and file_bytes.lstrip().startswith(_PDF_PREFIX):
        evidence_bytes = file_bytes
    if evidence_bytes:
        _retain_pdf_evidence(batch, evidence_bytes, signed_pdf_filename)

    expected_codes = set(accounts_by_code)

    received_sha256 = sha256(file_bytes).hexdigest()
    if received_sha256 != batch.content_sha256:
        message = (
            "Baseline workbook bytes do not match staged ImportBatch content_sha256 "
            f"(staged {batch.content_sha256}, received {received_sha256})."
        )
        report = _baseline_report(
            (),
            input_rows=0,
            errors=[message],
            expected_account_codes=expected_codes,
            allowed_account_codes=set(accounts_by_code),
        )
        fail_import_batch(session, batch, errors=[message], validation_report=report)
        raise BaselineValidationError(message)

    try:
        parsed = _parse_baseline_xlsx(
            file_bytes,
            expected_account_codes=expected_codes,
            allowed_account_codes=set(accounts_by_code),
            expected_currency=expected_currency,
            expected_balance_date=expected_balance_date,
        )
    except BaselineValidationError as error:
        report = _baseline_report(
            getattr(error, "records", ()),
            input_rows=getattr(error, "input_rows", 0),
            skipped_rows=getattr(error, "skipped_rows", 0),
            skip_reasons=getattr(error, "skip_reasons", ()),
            errors=[str(error)],
            expected_account_codes=expected_codes,
            allowed_account_codes=set(accounts_by_code),
        )
        fail_import_batch(session, batch, errors=[str(error)], validation_report=report)
        raise

    report = _baseline_report(
        parsed.records,
        input_rows=parsed.input_rows,
        skipped_rows=parsed.skipped_rows,
        skip_reasons=parsed.skip_reasons,
        expected_account_codes=expected_codes,
        allowed_account_codes=set(accounts_by_code),
    )
    for record in parsed.records:
        account = accounts_by_code[record.ledger_account_code]
        session.add(BalanceCheckpoint(
            import_batch_id=batch.id,
            entity_id=entity.id,
            ledger_account_id=account.id,
            balance_date=record.balance_date,
            balance=record.balance,
            currency=record.currency,
            source_metadata={
                **dict(record.source_metadata),
                "source_filename": batch.original_filename,
                "source_sha256": batch.content_sha256,
            },
        ))
    batch.source_metadata = {
        **(batch.source_metadata or {}),
        "parser": "xlsx_baseline",
        "required_headers": list(BASELINE_REQUIRED_HEADERS),
    }
    batch.validation_report = report
    session.flush()
    if activate:
        activate_import_batch(session, batch, validation_report=report)
    return report


parse_balance_checkpoint_xlsx = parse_baseline_xlsx
normalize_balance_checkpoint_xlsx = normalize_baseline_xlsx
ingest_baseline = normalize_baseline_xlsx
